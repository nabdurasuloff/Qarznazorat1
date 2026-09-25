# -*- coding: utf-8 -*-
"""
Word shabloniga ma'lumotlarni joylab, tayyor xat (.docx) yaratish.
"""
import os
import re
import sys
import datetime
from docx import Document
import util


def _base_dir():
    # PyInstaller --onefile bilan yig'ilganda fayllar vaqtinchalik papkaga
    # (sys._MEIPASS) ochiladi; oddiy ishga tushirishda esa shu faylning papkasi.
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        return sys._MEIPASS
    return os.path.dirname(os.path.abspath(__file__))


def _shablon_papkasi():
    """Foydalanuvchi yuklagan (moslashtirilgan) shablonlar saqlanadigan
    DOIMIY papka. Agar bu yerda moslashtirilgan shablon bo'lsa, u standart
    (dastur ichidagi) shablondan USTUN turadi.

    MUHIM: bu, backend.exe joylashgan papka o'rniga, foydalanuvchiga
    tegishli (har doim yoziladigan) papkada saqlanadi — chunki agar dastur
    "Program Files" kabi HIMOYALANGAN joyga o'rnatilgan bo'lsa, exe
    papkasiga yozish administratorsiz muvaffaqiyatsiz bo'lishi mumkin edi,
    va bu holatda shablon "yangilanmagandek" ko'rinardi."""
    if getattr(sys, 'frozen', False):
        asosiy = os.getenv('LOCALAPPDATA') or os.getenv('APPDATA') or os.path.expanduser('~')
        papka = os.path.join(asosiy, 'QarzNazorat', 'moslashtirilgan_shablonlar')
    else:
        papka = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'moslashtirilgan_shablonlar')
    os.makedirs(papka, exist_ok=True)
    return papka


def _effektiv_shablon_yoli(standart_yol, fayl_nomi):
    """Agar moslashtirilgan (foydalanuvchi yuklagan) shablon mavjud bo'lsa
    o'shani, aks holda dastur ichidagi standart shablonni qaytaradi."""
    moslashtirilgan = os.path.join(_shablon_papkasi(), fayl_nomi)
    return moslashtirilgan if os.path.exists(moslashtirilgan) else standart_yol


TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'xat_shablon.docx')
DAVO_ARIZA_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'davo_ariza_shablon.docx')
SUGURTA_XABARNOMA_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'sugurta_xabarnoma_shablon.docx')
YIGMA_JILD_TITUL_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'yigma_jild_titul_shablon.docx')
MALUMOTNOMA_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'malumotnoma_shablon.docx')
MALUMOTNOMA_TOPSHIRISHDA_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'malumotnoma_topshirishda_shablon.docx')
MALUMOTNOMA_KUN_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'malumotnoma_kun_shablon.docx')
SUGURTA_TOVON_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'sugurta_tovon_ariza_shablon.docx')
VAFOT_SUGURTA_TOVON_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'vafot_sugurta_tovon_shablon.docx')


def _fmt_summa(val):
    try:
        val = float(val)
    except (TypeError, ValueError):
        return str(val)
    return f"{val:,.0f}".replace(',', ' ')


def _replace_in_paragraph(paragraph, mapping):
    full_text = ''.join(run.text for run in paragraph.runs)
    if '{{' not in full_text:
        return
    new_text = full_text
    for key, val in mapping.items():
        # Katta-kichik harf farqiga sezgir bo'lmagan almashtirish — agar
        # shablon biror joyda (masalan Word orqali) qo'lda tahrirlanib,
        # placeholder harflari o'zgarib qolgan bo'lsa ham ishlashi uchun.
        pattern = re.compile(r'\{\{\s*' + re.escape(key) + r'\s*\}\}', re.IGNORECASE)
        new_text = pattern.sub(str(val), new_text)
    if new_text == full_text:
        return
    # Barcha runlarni tozalab, birinchi runga yangi matnni yozamiz
    # (formatlashni birinchi run’dan olamiz)
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(new_text)


def _replace_everywhere(doc, mapping):
    for p in doc.paragraphs:
        _replace_in_paragraph(p, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, mapping)


def generate_letter(output_path, xat_turi, mijoz_ism, mijoz_manzil, portfel_row, settings,
                     anketa_raqami=None, rahbar_ism=None):
    """
    xat_turi: 'Ogohlantirish' yoki 'Talabnoma'
    portfel_row: dict — bitta portfel qatoridagi ma'lumot (database.py formatida)
    settings: dict — get_all_settings() natijasi
    """
    doc = Document(_effektiv_shablon_yoli(TEMPLATE_PATH, 'xat_shablon.docx'))

    sarlavha = "OGOHLANTIRISH XATI" if xat_turi == 'Ogohlantirish' else "TALABNOMA"

    holat_sanasi = datetime.date.today().strftime('%d.%m.%Y')

    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'BANK_QISQA_NOMI': settings.get('bank_qisqa_nomi', ''),
        'BANK_MANZIL': settings.get('bank_manzil', ''),
        'BANK_EMAIL': settings.get('bank_email', ''),
        'BANK_SAYT': settings.get('bank_sayt', ''),
        'BANK_TEL': settings.get('bank_tel', ''),
        'BANK_MOBIL_ILOVA': settings.get('bank_mobil_ilova', ''),
        'BANK_KODI': settings.get('bank_kodi', ''),
        'ALOQA_MARKAZI_TEL': settings.get('aloqa_markazi_tel', ''),
        'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'FILIAL_TEL': settings.get('filial_tel', ''),
        'RAHBAR_ISM': rahbar_ism or settings.get('rahbar_ism', ''),

        'MIJOZ_ISM': mijoz_ism or '',
        'MIJOZ_MANZIL': mijoz_manzil or '',
        'SARLAVHA': sarlavha,

        'KREDIT_SUMMA': _fmt_summa(portfel_row.get('jami_berilgan_summa') or portfel_row.get('ead', 0)),
        'KREDIT_MAQSAD': portfel_row.get('tulov_maqsadi', '') or '',
        'HOLAT_SANASI': holat_sanasi,
        'JAMI_QARZ': _fmt_summa(portfel_row.get('jami_qarz', 0)),
        'ASOSIY_QARZ': _fmt_summa(portfel_row.get('asosiy_qarz', 0)),
        'FOIZ_QARZ': _fmt_summa(portfel_row.get('foiz_qarz', 0)),
        'JARIMA': _fmt_summa(portfel_row.get('jarima', 0)),
        'TOLOV_MUDDATI': f"{settings.get('tolov_muddati_kun', '10')} bank ish kuni",
        'ANKETA_RAQAM': anketa_raqami or portfel_row.get('anketa_raqami', ''),
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_davo_ariza(*args, **kwargs):
    raise NotImplementedError(
        "Bu funksiya eskirgan. O'rniga generate_davo_ariza_v2() dan foydalaning."
    )


def safe_filename(text):
    text = re.sub(r'[\\/*?:"<>|]', '', str(text))
    text = text.strip().replace(' ', '_')
    return text[:80]


DAVO_ARIZA_TEMPLATES = {
    'jismoniy_oddiy': 'davo_jismoniy_oddiy.docx',
    'jismoniy_kafil': 'davo_jismoniy_kafil.docx',
    'jismoniy_garov': 'davo_jismoniy_garov.docx',
    'jismoniy_kafil_garov': 'davo_jismoniy_kafil_garov.docx',
    'yuridik_oddiy': 'davo_yuridik_oddiy.docx',
    'yuridik_kafil': 'davo_yuridik_kafil.docx',
    'yuridik_kafil_garov': 'davo_yuridik_kafil_garov.docx',
    'yuridik_garov': 'davo_yuridik_garov.docx',
    'muddatidan_oldin': 'davo_muddatidan_oldin.docx',
}

DAVO_ARIZA_NOMLARI = {
    'jismoniy_oddiy': "Jismoniy shaxs — oddiy (kafilsiz)",
    'jismoniy_kafil': "Jismoniy shaxs — kafil bilan",
    'jismoniy_garov': "Jismoniy shaxs — faqat garov mulkiga qaratish (kafilsiz)",
    'jismoniy_kafil_garov': "Jismoniy shaxs — kafil + garov mulkiga qaratish",
    'yuridik_oddiy': "Yuridik/F.X — oddiy (kafilsiz)",
    'yuridik_kafil': "Yuridik shaxs — kafil bilan (foiz-penya)",
    'yuridik_kafil_garov': "Yuridik shaxs — kafil + garov mulkiga qaratish",
    'yuridik_garov': "Yuridik shaxs — faqat garov mulkiga qaratish (kafilsiz)",
    'muddatidan_oldin': "Shartnomani bekor qilish (muddatidan oldin undirish)",
}


def tavsiya_ariza_turi(mijoz_turi, taminot):
    """
    Mijoz turi va kiritilgan ta'minot (kafillik/garov) ma'lumotiga qarab,
    eng mos Davo ariza turini tavsiya qiladi.
    taminot: davo_taminot jadvalidagi dict (yoki None/bo'sh).
    """
    taminot_turi = (taminot or {}).get('taminot_turi') or 'yoq'
    # YaTT (yakka tartibdagi tadbirkor) — shaxsan jismoniy shaxs bo'lsa-da,
    # tadbirkorlik faoliyati yuritgani uchun Davo ariza turi tanlashda
    # yuridik shaxs kabi ko'riladi (iqtisodiy sud, tegishli shablon).
    if mijoz_turi in ('yuridik', 'yatt'):
        if taminot_turi == 'kafillik_garov':
            return 'yuridik_kafil_garov'
        elif taminot_turi == 'kafillik':
            return 'yuridik_kafil'
        elif taminot_turi == 'garov':
            return 'yuridik_garov'
        return 'yuridik_oddiy'
    else:
        if taminot_turi == 'kafillik_garov':
            return 'jismoniy_kafil_garov'
        elif taminot_turi == 'kafillik':
            return 'jismoniy_kafil'
        elif taminot_turi == 'garov':
            return 'jismoniy_garov'
        return 'jismoniy_oddiy'


def _oy_farqi(sana1_str, sana2_str):
    """Ikki sana orasidagi farqni oy hisobida qaytaradi (taxminiy)."""
    for fmt_pair in [('%d.%m.%Y', '%d.%m.%Y')]:
        try:
            d1 = datetime.datetime.strptime(str(sana1_str)[:10], '%d.%m.%Y')
            d2 = datetime.datetime.strptime(str(sana2_str)[:10], '%d.%m.%Y')
            months = (d2.year - d1.year) * 12 + (d2.month - d1.month)
            return max(months, 0)
        except (ValueError, TypeError):
            return ''
    return ''


def generate_davo_ariza_v2(turi, output_path, portfel_row, mijoz, taminot, settings,
                            xat_sanasi='', xat_turi_nomi='', imzo_ism_override=None):
    """
    turi: DAVO_ARIZA_TEMPLATES kalitlaridan biri.
    mijoz: mijozlar jadvalidagi dict (yoki None).
    taminot: davo_taminot jadvalidagi dict (yoki None) — kafil/garov ma'lumotlari.
    """
    tpl_file = DAVO_ARIZA_TEMPLATES.get(turi)
    if not tpl_file:
        raise ValueError(f"Noma'lum davo ariza turi: {turi}")
    tpl_path = _effektiv_shablon_yoli(os.path.join(_base_dir(), 'templates', tpl_file), tpl_file)
    if not os.path.exists(tpl_path):
        raise FileNotFoundError(
            f"Shablon fayli topilmadi: {tpl_file}\n\n"
            "Bu — dastur .exe qilib qayta yig'ilganda 'templates' papkasi "
            "eng so'nggi (to'liq) versiya bilan almashtirilmagani sababli "
            "yuz berishi mumkin. Dastur kodini eng so'nggi arxivdan qayta "
            "oching (templates papkasini to'liq almashtiring) va qaytadan "
            "build_exe.bat ishga tushiring."
        )
    doc = Document(tpl_path)

    mijoz = mijoz or {}
    taminot = taminot or {}
    holat_sanasi = datetime.date.today().strftime('%d.%m.%Y')
    muddat_oy = _oy_farqi(portfel_row.get('shartnoma_sanasi'), portfel_row.get('shartnoma_tugash_sanasi'))

    jami_ead = float(portfel_row.get('ead', 0) or 0)
    asosiy = float(portfel_row.get('asosiy_qarz', 0) or 0)
    foiz = float(portfel_row.get('foiz_qarz', 0) or 0)
    jarima_v = float(portfel_row.get('jarima', 0) or 0)
    jami = asosiy + foiz + jarima_v
    muddati_kelmagan = max(jami_ead - jami, 0)

    pochta = taminot.get('pochta_xarajati') or settings.get('pochta_xarajati_standart', '41200')

    mapping = {
        'ARIZA_SANA_QATORI': f'{holat_sanasi}-yil',
        'SUD_NOMI': settings.get('sud_iqtisodiy_nomi' if turi.startswith('yuridik') or turi == 'muddatidan_oldin'
                                  else 'sud_fuqarolik_nomi', ''),
        'PALATA_NOMI': settings.get('palata_nomi', ''),

        'BANK_NOMI': settings.get('bank_nomi', ''),
        'BANK_QISQA_NOMI': settings.get('bank_qisqa_nomi', ''),
        'BANK_STIR': settings.get('bank_stir', ''),
        'BANK_KODI_BOSH': settings.get('bank_kodi_bosh', ''),
        'BANK_HISOB_RAQAM_BOSH': settings.get('bank_hisob_raqami_bosh', ''),
        'BANK_MANZIL_BOSH': settings.get('bank_manzil', ''),
        'BANK_MANZIL_FILIAL': settings.get('bank_rasmiy_manzil_filial', ''),
        'BANK_REKVIZIT_QATORI': (f"ҳ|р: {settings.get('bank_hisob_raqami_filial', '')}, "
                                  f"банк коди: {settings.get('bank_kodi_filial', '')}, "
                                  f"СТИР: {settings.get('bank_stir', '')}, "
                                  f"Манзил: {settings.get('bank_rasmiy_manzil_filial', '')}"),
        'FILIAL_NOMI': settings.get('filial_nomi', ''),

        'MIJOZ_ISM': util.mijoz_ism_hujjat_uchun(
            mijoz.get('ism') or portfel_row.get('mijoz_nomi', ''),
            util.turi_kodidan(portfel_row.get('mijoz_turi_kodi'), portfel_row.get('mijoz_turi'))
        ),
        'MIJOZ_MANZIL': mijoz.get('manzil', '') or '',
        'MIJOZ_PINFL': portfel_row.get('pinfl', '') or mijoz.get('hujjat_raqami', '') or '',
        'MIJOZ_STIR': portfel_row.get('stir', '') or '',
        'MIJOZ_BANK_KODI': portfel_row.get('filial_kodi', '') or '',
        'MIJOZ_HISOB_RAQAM': portfel_row.get('kredit_hisob_raqami', '') or '',
        'MIJOZ_RAHBAR': mijoz.get('rahbar_ism', '') or '',
        'MIJOZ_TEL': mijoz.get('telefon', '') or '',
        'MIJOZ_PASSPORT_TOLIQ': _passport_toliq(mijoz),
        'BANK_HUDUDIY_NOMI': f"{settings.get('viloyat_nomi','Сирдарё')} вилоят худудий бошкармаси ({settings.get('filial_nomi','')} филиали)",

        'KAFIL_ISM': taminot.get('kafil_ism', '') or '',
        'KAFIL_MANZIL': taminot.get('kafil_manzil', '') or '',
        'KAFIL_PINFL': taminot.get('kafil_pinfl', '') or '',
        'KAFIL_TEL': taminot.get('kafil_tel', '') or '',
        'KAFIL_TUGILGAN_SANA': taminot.get('kafil_passport_sana', '') or '',
        'KAFIL_TUGILGAN_JOY': taminot.get('kafil_passport_organ', '') or '',
        'KAFIL_PASSPORT_TOLIQ': (
            f"{taminot.get('kafil_passport','')}, {taminot.get('kafil_passport_sana','')} йил "
            f"{taminot.get('kafil_passport_organ','')}дан берилган" if taminot.get('kafil_passport') else ''
        ),

        'GAROV_TAVSIFI': taminot.get('garov_tavsifi', '') or '',
        'GAROV_BAHOSI': _fmt_summa(taminot.get('garov_bahosi', 0)),

        'SHARTNOMA_SANA': portfel_row.get('shartnoma_sanasi', '') or '',
        'KREDIT_MUDDATI_OY': muddat_oy,
        'IMTIYOZLI_DAVR_OY': '',
        'YILLIK_FOIZ': portfel_row.get('yillik_foiz', '') or '',
        'KREDIT_MAQSAD': portfel_row.get('tulov_maqsadi', '') or '',
        'KREDIT_SUMMA': _fmt_summa(portfel_row.get('jami_berilgan_summa') or portfel_row.get('ead', 0)),

        'HOLAT_SANASI': holat_sanasi,
        'JAMI_QARZ': _fmt_summa(jami),
        'ASOSIY_QARZ': _fmt_summa(asosiy),
        'FOIZ_QARZ': _fmt_summa(foiz),
        'JARIMA': _fmt_summa(jarima_v),
        'MUDDATI_KELMAGAN_ASOSIY': _fmt_summa(muddati_kelmagan),
        'POCHTA_XARAJATI': _fmt_summa(pochta),

        'XAT_SANASI': xat_sanasi,
        'XAT_TURI_NOMI': xat_turi_nomi,
        'VOQEALAR_TAVSIFI': taminot.get('garov_tavsifi', '') or '[Voqealar tavsifini shu yerga kiriting]',

        'IMZO_ISM': imzo_ism_override or settings.get('sud_ariza_imzo_ism', '') or '',
        'IMZO_LAVOZIM': settings.get('sud_ariza_imzo_lavozimi', ''),
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def _passport_toliq(mijoz):
    if not mijoz:
        return ''
    hujjat = mijoz.get('hujjat_raqami', '') or ''
    sana = mijoz.get('passport_sana', '') or ''
    organ = mijoz.get('passport_organ', '') or ''
    if not hujjat:
        return ''
    parts = [hujjat]
    if sana:
        parts.append(f"{sana} йил")
    if organ:
        parts.append(f"{organ}дан берилган")
    return ', '.join(parts) if len(parts) > 1 else parts[0]


def _stdio_himoyasi():
    """PyInstaller --windowed rejimida sys.stdout/stderr None bo'lishi mumkin —
    docx2pdf shunga yozishga urinib xato beradi, shu sabab himoya qo'yamiz."""
    if sys.stdout is None:
        sys.stdout = open(os.devnull, 'w')
    if sys.stderr is None:
        sys.stderr = open(os.devnull, 'w')


# Kesh tozalash BIR MARTA bajarilganini eslab qolamiz (pastdagi izohga qarang).
_GEN_PY_TOZALANDI = False


def _gen_py_keshini_tozala():
    """docx2pdf ichida win32com.client.gencache.EnsureDispatch() ishlatiladi —
    bu "erta bog'lash" (early binding) uchun avtomatik generatsiya qilingan
    Python-COM ko'prik fayllarini keshga yozadi. PyInstaller bilan yig'ilgan
    (.exe) dasturda bu kesh ko'pincha yo'q yoki eskirgan bo'lib qoladi va
    "CLSIDToClass" kabi tushunarsiz xatoga olib keladi. Yechim — keshni
    tozalab, win32com'ni "sof" holatdan qayta generatsiya qilishga majburlash.

    MUHIM TEZLIK TUZATISHI: ilgari bu tozalash HAR BIR fayl uchun qayta-qayta
    bajarilardi. Natijada 8 ta hujjatli yig'ma jildda kesh 8 marta o'chirilib,
    8 marta qaytadan generatsiya qilinardi — jild tayyorlanishi daqiqalab
    cho'zilib ketardi. Aslida bu tozalash dastur ishga tushgandan keyin FAQAT
    BIR MARTA kerak, shuning uchun endi bayroq orqali bir martaga cheklandi.
    """
    global _GEN_PY_TOZALANDI
    if _GEN_PY_TOZALANDI:
        return
    _GEN_PY_TOZALANDI = True
    try:
        import win32com
        gen_py_path = os.path.join(win32com.__gen_path__)
        if os.path.isdir(gen_py_path):
            import shutil
            shutil.rmtree(gen_py_path, ignore_errors=True)
    except Exception:
        pass  # Bu tozalash muvaffaqiyatsiz bo'lsa ham, asosiy amal davom etadi


def convert_docx_to_pdf(docx_path, pdf_path=None, delete_docx=False):
    """
    .docx faylni .pdf ga aylantiradi. Windows'da MS Word o'rnatilgan bo'lishi
    shart (docx2pdf shu orqali ishlaydi). Agar Word bo'lmasa, xato chiqadi.
    """
    _stdio_himoyasi()
    _gen_py_keshini_tozala()

    if pdf_path is None:
        pdf_path = os.path.splitext(docx_path)[0] + '.pdf'
    try:
        from docx2pdf import convert
    except ImportError:
        raise RuntimeError(
            "PDF yaratish uchun 'docx2pdf' kutubxonasi o'rnatilmagan. "
            "requirements.txt orqali o'rnating: pip install docx2pdf pywin32"
        )
    try:
        convert(docx_path, pdf_path)
    except Exception as e:
        xabar = str(e)
        # Agar birinchi urinish keshga bog'liq sabab bilan muvaffaqiyatsiz
        # bo'lsa, "kech bog'lash" (late binding, Dispatch) usuli bilan
        # to'g'ridan-to'g'ri Word'ni boshqarib ko'ramiz — bu keshga
        # butunlay bog'liq emas va PyInstaller muhitida ancha ishonchli.
        try:
            _convert_docx_to_pdf_late_binding(docx_path, pdf_path)
        except Exception as e2:
            raise RuntimeError(
                f"PDF'ga aylantirishda xato: {xabar}\n"
                f"(muqobil usul ham muvaffaqiyatsiz: {e2})\n"
                "Bu funksiya faqat Windows'da, Microsoft Word o'rnatilgan bo'lsa ishlaydi. "
                "MS Word o'rnatilganini va boshqa hech qanday Word oynasi ochiq "
                "qolmaganini tekshiring."
            )
    if delete_docx and os.path.exists(docx_path):
        os.remove(docx_path)
    return pdf_path


def _convert_docx_to_pdf_late_binding(docx_path, pdf_path):
    """docx2pdf'ning standart (early-binding) usuli muvaffaqiyatsiz bo'lsa,
    Word'ni to'g'ridan-to'g'ri, keshga bog'liq bo'lmagan holda boshqarish."""
    import win32com.client
    import pythoncom
    pythoncom.CoInitialize()
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        # 17 = wdFormatPDF
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close()
        word.Quit()
    finally:
        pythoncom.CoUninitialize()


def convert_docx_to_pdf_batch(fayllar_royxati, chiqish_papkasi, xatolar=None):
    """Bir nechta Word hujjatini BITTA Word seansida PDF'ga aylantiradi.

    MUHIM TEZLIK TUZATISHI: ilgari yig'ma jild tayyorlanayotganda har bir
    hujjat uchun Word ALOHIDA ishga tushirilib, keyin yopilardi. Word'ning
    ishga tushishi o'zi 3-6 soniya olgani uchun, 8 hujjatli jild faqat
    Word'ni ochib-yopishga bir daqiqagacha vaqt sarflardi. Endi Word bir
    marta ochiladi, barcha hujjatlar ketma-ket aylantiriladi va oxirida
    bir marta yopiladi — jild bir necha barobar tez tayyor bo'ladi.

    Natija: {manba_fayl: pdf_fayl} lug'ati. Aylantirib bo'lmagan fayllar
    natijaga tushmaydi, ularning sababi 'xatolar' ro'yxatiga yoziladi.
    """
    natija = {}
    if not fayllar_royxati:
        return natija
    _stdio_himoyasi()
    _gen_py_keshini_tozala()
    os.makedirs(chiqish_papkasi, exist_ok=True)

    def _chiqish_yoli(idx, fayl):
        nom = safe_filename(os.path.splitext(os.path.basename(fayl))[0])[:60]
        return os.path.join(chiqish_papkasi, f"{idx:03d}_{nom}.pdf")

    # 1-usul: bitta Word seansi (eng tez)
    try:
        import pythoncom
        import win32com.client
    except Exception:
        pythoncom = None

    if pythoncom is not None:
        word = None
        try:
            pythoncom.CoInitialize()
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            try:
                word.DisplayAlerts = 0
            except Exception:
                pass
            for idx, fayl in enumerate(fayllar_royxati, 1):
                chiqish = _chiqish_yoli(idx, fayl)
                try:
                    doc = word.Documents.Open(os.path.abspath(fayl), ReadOnly=True)
                    try:
                        doc.SaveAs(os.path.abspath(chiqish), FileFormat=17)  # 17 = wdFormatPDF
                    finally:
                        doc.Close(False)
                    natija[fayl] = chiqish
                except Exception:
                    # Bu faylni keyin eski (bir faylli) usul bilan sinab ko'ramiz
                    pass
        except Exception:
            pass  # Word umuman ochilmadi — pastdagi zaxira usulga o'tamiz
        finally:
            if word is not None:
                try:
                    word.Quit()
                except Exception:
                    pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    # 2-usul (zaxira): qolgan fayllarni eski, bir faylli yo'l bilan
    for idx, fayl in enumerate(fayllar_royxati, 1):
        if fayl in natija:
            continue
        chiqish = _chiqish_yoli(idx, fayl)
        try:
            convert_docx_to_pdf(fayl, chiqish, delete_docx=False)
            natija[fayl] = chiqish
        except Exception as e:
            if xatolar is not None:
                xatolar.append(f"{os.path.basename(fayl)}: PDF'ga aylantirib bo'lmadi ({e})")
    return natija


REESTR_SSP_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'reestr_ssp_shablon.docx')


def _paragraf_almashtir(p, eski, yangi):
    """Paragraf ichidagi matnni (bir nechta run'ga bo'lingan bo'lsa ham) to'g'ri almashtiradi."""
    toliq = ''.join(r.text for r in p.runs)
    if eski not in toliq:
        return False
    yangi_toliq = toliq.replace(eski, yangi)
    if p.runs:
        p.runs[0].text = yangi_toliq
        for r in p.runs[1:]:
            r.text = ''
    return True


def _row_katakcha_nusxala(manba_row, yangi_row):
    """Bitta jadval qatoridagi formatlashni (shrift, chegara) boshqa qatorga nusxalaydi."""
    import copy
    for i, manba_cell in enumerate(manba_row.cells):
        if i >= len(yangi_row.cells):
            break
        yangi_cell = yangi_row.cells[i]
        manba_tcPr = manba_cell._tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
        if manba_tcPr is not None:
            yangi_tcPr = yangi_cell._tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
            if yangi_tcPr is not None:
                yangi_cell._tc.remove(yangi_tcPr)
            yangi_cell._tc.insert(0, copy.deepcopy(manba_tcPr))


def _birlashtir_docx_xom(fayllar_royxati, output_path):
    """
    XOM (past darajadagi) Word birlashtirish — faqat bir XIL shablon
    turidagi hujjatlar uchun xavfsiz ishlaydi (XML tuzilmasi bir xil
    bo'lgani uchun). Turli shablon (masalan jismoniy_oddiy + yuridik_kafil)
    aralashtirilsa, ichki tuzilma farqlari kontent noto'g'ri joylashib
    qolishiga olib kelishi mumkin — shu sabab bu funksiya faqat
    birlashtir_hujjatlar() ichida, xavfsiz ekani tasdiqlangan holatlardagina
    chaqiriladi.

    MUHIM: bu funksiya faqat python-docx'ning o'zidan (XML elementlarini
    to'g'ridan-to'g'ri nusxalash orqali) foydalanadi — tashqi
    'docxcompose' kutubxonasiga BOG'LIQ EMAS. Sabab: docxcompose o'zining
    ichki shablon fayllarini (templates/*.xml) talab qiladi, va bu
    fayllar PyInstaller bilan .exe qilib yig'ilganda ba'zan to'g'ri
    qadoqlanmay, "No such file" xatosiga olib kelishi mumkin edi. Bu
    yechim esa hech qanday qo'shimcha faylga muhtoj emas — har doim
    ishlaydi.

    fayllar_royxati: fayl yo'llari ro'yxati (kamida bitta bo'lishi kerak).
    """
    import copy as _copy

    if not fayllar_royxati:
        raise ValueError("Birlashtirish uchun kamida bitta fayl kerak.")

    asosiy = Document(fayllar_royxati[0])

    def _paragraf_boshmi(element):
        """Elementda matn yoki rasm bor-yo'qligini tekshiradi (haqiqiy
        mazmunga ega birinchi elementni topish uchun) — nomlar maydoni
        (namespace) e'lonlaridagi 'drawing' so'ziga aldanmaslik uchun,
        faqat w:t (matn) va w:drawing/w:pict teglarining o'zini tekshiradi."""
        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        for t_el in element.iter(f'{ns}t'):
            if t_el.text and t_el.text.strip():
                return True
        for tag in ('drawing', 'pict', 'tbl'):
            if element.tag.endswith('}' + tag):
                return True
            if list(element.iter(f'{{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}}{tag}')) or \
               list(element.iter(f'{ns}{tag}')):
                return True
        return False

    for fayl in fayllar_royxati[1:]:
        qoshimcha = Document(fayl)
        elements = list(qoshimcha.element.body)
        # Boshidagi to'liq BO'SH paragraflarni o'tkazib yuboramiz — aks
        # holda ular yangi sahifada foydasiz bo'sh joy hosil qilib
        # qo'yishi mumkin edi.
        boshlanish = 0
        for idx, el in enumerate(elements):
            if el.tag.endswith('}sectPr'):
                continue
            if _paragraf_boshmi(el):
                boshlanish = idx
                break
        else:
            boshlanish = 0

        birinchi = True
        for element in elements[boshlanish:]:
            if element.tag.endswith('}sectPr'):
                continue
            nusxa = _copy.deepcopy(element)
            ns_local = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            if nusxa.tag.endswith('}p'):
                pPr = nusxa.find(f'{ns_local}pPr')
                ichki_sectpr = pPr.find(f'{ns_local}sectPr') if pPr is not None else None
                if ichki_sectpr is not None:
                    pPr.remove(ichki_sectpr)
                if birinchi:
                    # MUHIM: alohida "sahifa o'tkazish" paragrafi qo'shish
                    # o'rniga (bu ba'zan LibreOffice/Word'da keraksiz bo'sh
                    # sahifa hosil qilib qo'yadi), shu hujjatning ENG
                    # BIRINCHI paragrafiga bevosita "pageBreakBefore"
                    # xossasini beramiz — bu ancha ishonchli usul.
                    if pPr is None:
                        from docx.oxml.ns import qn
                        pPr = nusxa.makeelement(qn('w:pPr'), {})
                        nusxa.insert(0, pPr)
                    if pPr.find(f'{ns_local}pageBreakBefore') is None:
                        from docx.oxml.ns import qn
                        pbb = nusxa.makeelement(qn('w:pageBreakBefore'), {})
                        pPr.insert(0, pbb)
                    birinchi = False
            asosiy.element.body.append(nusxa)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    asosiy.save(output_path)
    return output_path


WORD_KENGAYTMALARI = ('.docx', '.doc', '.rtf', '.odt')
RASM_KENGAYTMALARI = ('.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff', '.gif', '.webp')


def _rasmni_pdf_qilish(rasm_yoli, pdf_yoli):
    """Skaner qilingan rasmni (JPG/PNG/TIFF) bir sahifali PDF'ga aylantiradi.

    MUHIM: mijozning kredit shartnomasi, garov hujjati va boshqa qo'lda
    yuklangan hujjatlar ko'pincha telefonda suratga olingan JPG/PNG bo'ladi.
    Ilgari bunday fayl Word hujjati deb hisoblanib, PDF'ga aylantirishda
    xato berardi va BUTUN yig'ma jild tayyorlanmay qolardi.
    """
    from PIL import Image
    img = Image.open(rasm_yoli)
    try:
        if img.mode in ('RGBA', 'LA', 'P'):
            img = img.convert('RGB')
        elif img.mode not in ('RGB', 'L', 'CMYK'):
            img = img.convert('RGB')
        img.save(pdf_yoli, 'PDF', resolution=150.0)
    finally:
        try:
            img.close()
        except Exception:
            pass
    return pdf_yoli


def _birlashtir_pdf(fayllar_royxati, output_path, xatolar=None):
    """
    Har bir hujjatni ALOHIDA PDF'ga aylantirib, so'ng tayyor PDF
    sahifalarini qo'shib (pypdf orqali) birlashtiradi. Bu — turli shablon
    turlari (masalan jismoniy_oddiy + yuridik_kafil) aralash bo'lganda
    ham KONTENT HECH QACHON ARALASHIB QOLMASLIGINI kafolatlaydi, chunki
    har bir hujjat avval MUSTAQIL, TO'LIQ tayyor holatda PDF qilinadi —
    keyin esa faqat tayyor PDF SAHIFALARI (rasmga o'xshash, tuzilmasi
    aralashmaydigan) qo'shiladi.

    'xatolar' — ro'yxat berilsa, jildga QO'SHILMAGAN hujjatlar va ularning
    sababi shu ro'yxatga yoziladi (chaqiruvchi taraf buni foydalanuvchiga
    ogohlantirish sifatida ko'rsatadi). Bitta hujjat buzuq bo'lgani uchun
    butun jild tayyorlanmay qolmasligi kerak.
    """
    import pypdf
    import shutil
    import tempfile

    if xatolar is None:
        xatolar = []

    # 1-BOSQICH. Ro'yxatni tozalash: bo'sh, yo'q va TAKRORLANGAN fayllar.
    # MUHIM: ilgari bir xil fayl ikki maydonda yozilgan bo'lsa (masalan
    # 'fayl_yoli' va 'davo_ariza_fayl_yoli' bir xil hujjatga ishora qilsa),
    # u jildga IKKI MARTA tushib qolardi.
    tozalangan, korilgan = [], set()
    for fayl in (fayllar_royxati or []):
        if not fayl:
            continue
        if not os.path.exists(fayl):
            xatolar.append(f"{os.path.basename(fayl)}: fayl topilmadi")
            continue
        kalit = os.path.normcase(os.path.realpath(fayl))
        if kalit in korilgan:
            continue
        korilgan.add(kalit)
        tozalangan.append(fayl)

    if not tozalangan:
        raise ValueError("Birlashtirish uchun kamida bitta mavjud fayl kerak.")

    # MUHIM TUZATISH (fayl yo'qolishi): ilgari .docx yonidagi <shu nom>.pdf
    # faylga yozilar, keyin esa u O'CHIRILARDI. Agar papkada hujjatning
    # imzolangan skaneri ham <shu nom>.pdf bo'lib tursa — u ustidan yozilib,
    # so'ng butunlay o'chib ketardi. Endi barcha oraliq PDF'lar alohida
    # vaqtinchalik papkada yaratiladi va mijoz papkasiga umuman tegilmaydi.
    vaqt_papka = tempfile.mkdtemp(prefix='yigma_jild_')
    try:
        word_fayllar = [f for f in tozalangan if f.lower().endswith(WORD_KENGAYTMALARI)]
        pdf_xarita = convert_docx_to_pdf_batch(word_fayllar, vaqt_papka, xatolar) if word_fayllar else {}

        writer = pypdf.PdfWriter()
        # MUHIM: PdfReader sahifalarni "dangasa" (lazy) o'qiydi — faylni
        # yozish paytigacha oqim ochiq turishi kerak. Shu sabab barcha
        # o'quvchilarni ro'yxatda ushlab turamiz, aks holda ular vaqtidan
        # oldin yopilib "seek of closed file" xatosi chiqishi mumkin edi.
        oquvchilar = []
        for idx, fayl in enumerate(tozalangan, 1):
            past = fayl.lower()
            try:
                if past.endswith('.pdf'):
                    pdf_yoli = fayl
                elif past.endswith(WORD_KENGAYTMALARI):
                    pdf_yoli = pdf_xarita.get(fayl)
                    if not pdf_yoli:
                        continue  # sababi allaqachon xatolar ro'yxatida
                elif past.endswith(RASM_KENGAYTMALARI):
                    pdf_yoli = _rasmni_pdf_qilish(
                        fayl, os.path.join(vaqt_papka, f"{idx:03d}_rasm.pdf"))
                else:
                    xatolar.append(
                        f"{os.path.basename(fayl)}: bu turdagi fayl jildga qo'shilmaydi "
                        f"(faqat PDF, Word va rasm fayllari qo'shiladi)")
                    continue

                reader = pypdf.PdfReader(pdf_yoli)
                if reader.is_encrypted:
                    # Parolsiz himoyalangan PDF'ni ochishga urinib ko'ramiz
                    try:
                        reader.decrypt('')
                    except Exception:
                        pass
                sahifalar = list(reader.pages)
                oquvchilar.append(reader)
                for page in sahifalar:
                    writer.add_page(page)
            except Exception as e:
                xatolar.append(f"{os.path.basename(fayl)}: {e}")

        if len(writer.pages) == 0:
            raise RuntimeError(
                "Hech bir hujjatni jildga qo'shib bo'lmadi:\n• " + "\n• ".join(xatolar or ['sabab aniqlanmadi']))

        papka = os.path.dirname(os.path.abspath(output_path))
        if papka:
            os.makedirs(papka, exist_ok=True)
        with open(output_path, 'wb') as f:
            writer.write(f)
    finally:
        shutil.rmtree(vaqt_papka, ignore_errors=True)
    return output_path


def birlashtir_hujjatlar(fayllar_royxati, output_path, ariza_turlari=None):
    """
    Bir nechta hujjatni (Davo arizalarni) BITTA faylga jamlaydi — natija
    output_path'ning kengaytmasiga qarab .docx yoki .pdf bo'ladi.

    ariza_turlari: har bir fayl uchun mos keluvchi Davo ariza turi kaliti
    ro'yxati (masalan ['jismoniy_oddiy', 'jismoniy_oddiy', 'yuridik_kafil']).
    Agar BARCHA fayllar BIR XIL turdan bo'lsa — tezkor, tahrirlash mumkin
    bo'lgan Word darajasidagi birlashtirish ishlatiladi (xavfsiz, chunki
    tuzilma bir xil). Agar turlar ARALASH bo'lsa (yoki ariza_turlari
    berilmagan bo'lsa) — har doim ishonchli PDF-orqali birlashtirish
    ishlatiladi, chunki turli shablonlar ichki tuzilmasi farq qilishi
    kontentning noto'g'ri joylashib qolishiga olib kelishi mumkin edi.
    """
    bir_xil_turdami = ariza_turlari is not None and len(set(ariza_turlari)) <= 1

    if output_path.lower().endswith('.docx') and bir_xil_turdami:
        return _birlashtir_docx_xom(fayllar_royxati, output_path)

    # PDF natija so'ralgan BO'LSA, yoki turlar aralash bo'lsa — PDF orqali
    if output_path.lower().endswith('.docx'):
        # Turlar aralash, lekin foydalanuvchi Word so'ragan — PDF qilib,
        # so'ng nomini .docx'ga o'zgartirmasdan, aniq PDF sifatida qaytaramiz
        # (chaqiruvchi taraf bu holatni alohida ko'rsatishi kerak).
        pdf_path = os.path.splitext(output_path)[0] + '.pdf'
        return _birlashtir_pdf(fayllar_royxati, pdf_path)

    return _birlashtir_pdf(fayllar_royxati, output_path)


def generate_reestr_ssp(output_path, mijozlar_royxati, xat_raqami, xat_sanasi, settings):
    """
    SSPga (Savdo-Sanoat Palatasiga) yuboriladigan Davo arizalar reestri
    (ro'yxat xati) yaratadi — bank rasmiy blankasi asosida, har bir
    tayyorlangan Davo ariza uchun anketa raqami, F.I.Sh va Davo summasi
    bilan jadval qator-qator to'ldiriladi, oxirida umumiy summa chiqariladi.

    mijozlar_royxati: har biri {'anketa_raqami', 'mijoz_ism', 'summa'} dict
    ro'yxati (summa — davo arizada ko'rsatilgan, sudga topshirilayotgan summa).
    """
    tpl_path = _effektiv_shablon_yoli(REESTR_SSP_TEMPLATE_PATH, 'reestr_ssp_shablon.docx')
    if not os.path.exists(tpl_path):
        raise FileNotFoundError(
            "Shablon fayli topilmadi: reestr_ssp_shablon.docx\n\n"
            "Bu — dastur .exe qilib qayta yig'ilganda 'templates' papkasi "
            "eng so'nggi (to'liq) versiya bilan almashtirilmagani sababli "
            "yuz berishi mumkin. Dastur kodini eng so'nggi arxivdan qayta "
            "oching (templates papkasini to'liq almashtiring) va qaytadan "
            "build_exe.bat ishga tushiring."
        )
    doc = Document(tpl_path)

    for p in doc.paragraphs:
        _paragraf_almashtir(p, '{{XAT_SANASI}}', xat_sanasi or '')
        _paragraf_almashtir(p, '{{XAT_RAQAMI}}', xat_raqami or '')
        _paragraf_almashtir(p, '{{IMZO_ISM}}', settings.get('sud_ariza_imzo_ism', '') or '')

    t = doc.tables[0]
    jami_row = t.rows[-1]  # "Жами" qatori - eng oxirida qolishi kerak

    jami_summa = 0.0
    for idx, m in enumerate(mijozlar_royxati, start=1):
        summa = m.get('summa') or 0
        jami_summa += summa
        yangi_row = t.add_row()
        _row_katakcha_nusxala(t.rows[0], yangi_row)
        yangi_row.cells[0].text = str(idx)
        yangi_row.cells[1].text = str(m.get('anketa_raqami', ''))
        yangi_row.cells[2].text = str(m.get('mijoz_ism', ''))
        yangi_row.cells[3].text = f"{summa:,.1f}".replace(',', ' ')

    # "Жами" qatorini jadvalning eng oxiriga ko'chiramiz (add_row() yangi
    # qatorlarni oxiriga qo'shgani uchun, Жами hozir o'rtada qolgan bo'ladi)
    tbl = t._tbl
    tbl.remove(jami_row._tr)
    tbl.append(jami_row._tr)
    _paragraf_almashtir(jami_row.cells[3].paragraphs[0], '{{JAMI_SUMMA}}',
                         f"{jami_summa:,.1f}".replace(',', ' '))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_dpd_kamaygan_yakunlash_hujjati(output_path, xat, portfel_row, joriy_dpd, eski_qarz, joriy_qarz):
    """Davo ariza tayyor (SSPdan qaytgan), lekin hali sudga topshirilmagan
    holatda — mijozning DPD/qarzi chegaradan pastga tushib qolganda,
    tizim ishni AVTOMATIK yakunlaydi. Bu hujjat — shu yakunlashning
    ASOSI sifatida, qarzdorlikning o'zgarishini ko'rsatadi."""
    doc = Document()
    doc.add_heading("Sud ishini DPD/qarz kamaygani sababli avtomatik yakunlash to'g'risida ma'lumotnoma", level=1)
    p = doc.add_paragraph()
    p.add_run("Mijoz: ").bold = True
    p.add_run(xat.get('mijoz_nomi', ''))
    p2 = doc.add_paragraph()
    p2.add_run("Anketa raqami: ").bold = True
    p2.add_run(str(xat.get('anketa_raqami', '')))
    p3 = doc.add_paragraph()
    p3.add_run("Davo ariza ish raqami: ").bold = True
    p3.add_run(str(xat.get('davo_ariza_ish_raqami', '') or '—'))
    p4 = doc.add_paragraph()
    p4.add_run("Hozirgi DPD (muddati o'tgan kun soni): ").bold = True
    p4.add_run(str(joriy_dpd))

    doc.add_paragraph()
    doc.add_heading("Qarzdorlikning o'zgarishi", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = '', "Qarzdorlik (so'm)"
    r1 = table.add_row().cells
    r1[0].text = 'Davo ariza summasi (sudga topshirish uchun tayyorlangan)'
    r1[1].text = f"{eski_qarz:,.0f}".replace(',', ' ')
    r2 = table.add_row().cells
    r2[0].text = 'Hozirgi kunda (joriy portfel)'
    r2[1].text = f"{joriy_qarz:,.0f}".replace(',', ' ')
    r3 = table.add_row().cells
    r3[0].text = "Farq"
    farq = joriy_qarz - eski_qarz
    belgi = '+' if farq >= 0 else ''
    r3[1].text = f"{belgi}{farq:,.0f}".replace(',', ' ')

    doc.add_paragraph()
    doc.add_paragraph(
        "Yuqoridagi ko'rsatkichlar asosida, mijozning qarzdorlik holati sudga "
        "topshirish zarurati qolmaydigan darajada yaxshilangani aniqlandi. "
        "Shu sababli, ushbu Davo ariza sudga topshirilmasdan, ish avtomatik "
        "yakunlandi. Agar keyinchalik qarzdorlik holati yana yomonlashsa, "
        "jarayon yangidan (nollashtirilgan holda) boshlanishi mumkin."
    )
    doc.add_paragraph()
    doc.add_paragraph(f"Sana: {datetime.date.today().strftime('%d.%m.%Y')}")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path



def generate_tolov_asosida_yakunlash_hujjati(output_path, xat, mijoz_nomi, jami_qarz, tolovlar_royxati, jami_tolangan):
    """MIB ishi to'lovlar orqali TO'LIQ to'langanda, tizim AVTOMATIK
    yakunlaganda — asos sifatida ko'rsatiladigan hujjatni yaratadi. Bu
    hujjat qaysi to'lovlar, qachon, qancha summada kelib tushgani va
    ularning yig'indisi qarzni to'liq qoplaganini ko'rsatadi."""
    doc = Document()
    doc.add_heading("MIB ishini to'lovlar asosida yakunlash to'g'risida ma'lumotnoma", level=1)
    p = doc.add_paragraph()
    p.add_run(f"Mijoz: ").bold = True
    p.add_run(mijoz_nomi)
    p2 = doc.add_paragraph()
    p2.add_run(f"Anketa raqami: ").bold = True
    p2.add_run(str(xat.get('anketa_raqami', '')))
    p3 = doc.add_paragraph()
    p3.add_run(f"MIB ish raqami: ").bold = True
    p3.add_run(str(xat.get('mib_ish_raqami', '') or '—'))
    p4 = doc.add_paragraph()
    p4.add_run(f"Davo (ijro) summasi: ").bold = True
    p4.add_run(f"{jami_qarz:,.0f} so'm".replace(',', ' '))

    doc.add_heading("Amalga oshirilgan to'lovlar ro'yxati", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Sana'
    hdr[1].text = 'Summa'
    hdr[2].text = 'Manba'
    hdr[3].text = 'PINFL/STIR/Hisob'
    for t in tolovlar_royxati:
        row = table.add_row().cells
        row[0].text = str(t.get('sana', ''))
        row[1].text = f"{(t.get('summa') or 0):,.0f}".replace(',', ' ')
        row[2].text = "Kunlik (29801)" if t.get('manba') == 'kunlik_29801' else "MIBdan"
        row[3].text = str(t.get('pinfl_yoki_stir', ''))

    doc.add_paragraph()
    yakun = doc.add_paragraph()
    yakun.add_run(f"Jami to'langan: ").bold = True
    yakun.add_run(f"{jami_tolangan:,.0f} so'm".replace(',', ' '))
    xulosa = doc.add_paragraph()
    xulosa.add_run(
        "Xulosa: yuqoridagi to'lovlar yig'indisi mijozning davo (ijro) summasini to'liq qoplagani "
        "uchun, MIB ishi tizim tomonidan AVTOMATIK ravishda yakunlangan deb belgilandi."
    ).italic = True
    doc.add_paragraph(f"Hujjat yaratilgan sana: {datetime.date.today().strftime('%d.%m.%Y')}")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_malumotnoma_topshirishda(output_path, xat, portfel_row, mijoz, settings):
    """Sudga topshirish paytidagi qarzdorlik ma'lumotnomasi. Bu — Davo ariza
    summasiga TENG bo'lishi kerak (sog'lom tekshiruv sifatida), chunki bir xil
    kundagi holatni aks ettiradi."""
    tpl = _effektiv_shablon_yoli(MALUMOTNOMA_TOPSHIRISHDA_TEMPLATE_PATH, 'malumotnoma_topshirishda_shablon.docx')
    doc = Document(tpl)
    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))
    portfel_row = portfel_row or {}
    pinfl_stir = portfel_row.get('pinfl') or portfel_row.get('stir') or ''
    # MUHIM: bu ma'lumotnoma — Davo ariza SUDGA TOPSHIRILAYOTGAN paytdagi
    # holatni aks ettiradi, shuning uchun qarzdorlik raqamlari albatta
    # Davo ariza summasiga (xat.davo_summasi_*) TENG bo'lishi shart —
    # BUGUNGI (joriy, o'zgargan bo'lishi mumkin) portfel qiymati emas.
    # Aks holda, agar bu hujjat Davo ariza yaratilgandan keyin (masalan
    # bir necha kun o'tib) tuzilsa, raqamlar ikkalasida FARQ qilib qolib,
    # "nega bir xil emas" degan savol tug'dirar edi.
    asosiy = xat.get('davo_summasi_asosiy') or 0
    foiz = xat.get('davo_summasi_foiz') or 0
    penya = xat.get('davo_summasi_jarima') or 0
    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''), 'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'QARZDOR_ISM': mijoz_ism, 'PINFL_STIR': pinfl_stir or '—',
        'SANA': datetime.date.today().strftime('%d.%m.%Y'),
        'ANKETA_RAQAM': xat.get('anketa_raqami', ''), 'UNIKAL_RAQAM': portfel_row.get('unikal') or '—',
        'KREDIT_MAQSADI': portfel_row.get('tulov_maqsadi') or '—',
        'KREDIT_SUMMA': _fmt_summa(portfel_row.get('jami_berilgan_summa') or 0),
        'YILLIK_FOIZ': portfel_row.get('yillik_foiz') or '—',
        'SHARTNOMA_SANASI': portfel_row.get('shartnoma_sanasi') or '—',
        'SHARTNOMA_TUGASH_SANASI': portfel_row.get('shartnoma_tugash_sanasi') or '—',
        # MUHIM: "jami asosiy qarz" — bu, umumiy KREDIT QOLDIG'I (hali
        # muddati kelmagan qismini ham o'z ichiga oladi) — joriy portfeldan
        # olinadi. "Muddati o'tgan" qismlar esa — aynan Davo arizada
        # SUD SO'RAYOTGAN summa bilan TENG bo'lishi shart, shuning uchun
        # ular joriy portfeldan emas, Davo ariza summasidan olinadi:
        'JAMI_ASOSIY_QARZ': _fmt_summa(portfel_row.get('jami_qarz') or (asosiy + foiz + penya)),
        'MUDDATI_OTGAN_ASOSIY': _fmt_summa(asosiy),
        'DPD_ASOSIY': portfel_row.get('dpd_asosiy') or 0,
        'JAMI_FOIZ': _fmt_summa(foiz),
        'DPD_FOIZ': portfel_row.get('dpd_foiz') or 0,
        'JAMI_PENYA': _fmt_summa(penya),
        'BOSHQARUVCHI_ISM': settings.get('boshqaruvchi_ism', '') or '—',
        # Eski, oddiy (jami) qiymatlar — orqaga moslik uchun saqlanadi
        'JAMI_QARZ': _fmt_summa(asosiy + foiz + penya),
        'ASOSIY_QARZ': _fmt_summa(asosiy), 'FOIZ_QARZ': _fmt_summa(foiz), 'PENYA_QARZ': _fmt_summa(penya),
    }
    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_malumotnoma_sud_kuni(output_path, xat, portfel_row, mijoz, settings):
    """Sud kunidan 1 kun oldin tayyorlanadigan ma'lumotnoma — sudga
    topshirilgan paytdagi (Davo summasi) va HOZIRGI (joriy portfel)
    qarzdorlikni solishtirib, farqni (asosiy/foiz/penya bo'yicha alohida)
    ko'rsatadi."""
    tpl = _effektiv_shablon_yoli(MALUMOTNOMA_KUN_TEMPLATE_PATH, 'malumotnoma_kun_shablon.docx')
    doc = Document(tpl)
    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))
    pinfl_stir = (portfel_row.get('pinfl') or portfel_row.get('stir') or '') if portfel_row else ''

    eski_asosiy = xat.get('davo_summasi_asosiy') or 0
    eski_foiz = xat.get('davo_summasi_foiz') or 0
    eski_penya = xat.get('davo_summasi_jarima') or 0
    joriy_asosiy = (portfel_row.get('asosiy_qarz') or 0) if portfel_row else 0
    joriy_foiz = (portfel_row.get('foiz_qarz') or 0) if portfel_row else 0
    joriy_penya = (portfel_row.get('jarima') or 0) if portfel_row else 0

    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''), 'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'QARZDOR_ISM': mijoz_ism, 'PINFL_STIR': pinfl_stir or '—',
        'SANA': datetime.date.today().strftime('%d.%m.%Y'),
        'ANKETA_RAQAM': xat.get('anketa_raqami', ''), 'UNIKAL_RAQAM': (portfel_row.get('unikal') or '—') if portfel_row else '—',
        'JAMI_QARZ': _fmt_summa(joriy_asosiy + joriy_foiz + joriy_penya),
        'ASOSIY_QARZ': _fmt_summa(joriy_asosiy), 'FOIZ_QARZ': _fmt_summa(joriy_foiz), 'PENYA_QARZ': _fmt_summa(joriy_penya),
    }
    _replace_everywhere(doc, mapping)

    doc.add_paragraph()
    doc.add_heading("Sudga topshirilgan kundagi qarzdorlik bilan solishtirish", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = '', 'Asosiy qarz', 'Foiz', 'Penya'
    r1 = table.add_row().cells
    r1[0].text = 'Sudga topshirilganda'
    r1[1].text, r1[2].text, r1[3].text = _fmt_summa(eski_asosiy), _fmt_summa(eski_foiz), _fmt_summa(eski_penya)
    r2 = table.add_row().cells
    r2[0].text = 'Hozirgi kunda'
    r2[1].text, r2[2].text, r2[3].text = _fmt_summa(joriy_asosiy), _fmt_summa(joriy_foiz), _fmt_summa(joriy_penya)
    r3 = table.add_row().cells
    r3[0].text = 'Farq (+o\'sish / -kamayish)'
    for i, (yangi, eski) in enumerate([(joriy_asosiy, eski_asosiy), (joriy_foiz, eski_foiz), (joriy_penya, eski_penya)]):
        farq = yangi - eski
        belgi = '+' if farq >= 0 else ''
        r3[i + 1].text = f"{belgi}{_fmt_summa(farq)}"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


SUD_YIGMA_JILD_TITUL_TEMPLATE_PATH = os.path.join(_base_dir(), 'templates', 'sud_yigma_jild_titul_shablon.docx')


def generate_sud_yigma_jild_titul(output_path, xat, portfel_row, mijoz, settings):
    """Sud harakatlari yig'ma jildi uchun titul (muqova) hujjati — MIB
    yig'ma jildi titulining AYNAN O'ZI (bir xil tuzilish, bir xil
    maydonlar), faqat o'z alohida shabloniga ega."""
    doc = Document(_effektiv_shablon_yoli(SUD_YIGMA_JILD_TITUL_TEMPLATE_PATH, 'sud_yigma_jild_titul_shablon.docx'))

    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '') or \
        (portfel_row.get('mijoz_nomi', '') if portfel_row else '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))
    mijoz_manzil = (mijoz.get('manzil') if mijoz else '') or ''
    hujjat_raqami = (mijoz.get('hujjat_raqami') if mijoz else '') or ''
    jami_qarz = ((portfel_row.get('asosiy_qarz', 0) or 0) + (portfel_row.get('foiz_qarz', 0) or 0) +
                 (portfel_row.get('jarima', 0) or 0)) if portfel_row else 0

    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'MIB_ISH_RAQAMI': xat.get('mib_ish_raqami', '') or '—',
        'QARZDOR_ISM': mijoz_ism,
        'MIJOZ_TURI': {'yuridik': "Yuridik shaxs", 'yatt': "Yakka tartibdagi tadbirkor (YaTT)"}.get(
            xat.get('mijoz_turi'), "Jismoniy shaxs"),
        'MIJOZ_MANZIL': mijoz_manzil or '—',
        'HUJJAT_RAQAMI': hujjat_raqami or '—',
        'ANKETA_RAQAM': xat.get('anketa_raqami', '') or '—',
        'JAMI_QARZ': f"{jami_qarz:,.0f}".replace(',', ' '),
        'XAT_SANASI': _fmt_sana_iso(xat.get('yaratilgan_sana')) or '—',
        'XAT_YUBORILGAN_SANA': _fmt_sana_iso(xat.get('yuborilgan_sana')) or '—',
        'DAVO_ARIZA_SANA': _fmt_sana_iso(xat.get('davo_ariza_sana')) or '—',
        'PALATADAN_QAYTGAN_SANA': xat.get('davo_ariza_imzo_sana', '') or '—',
        'SUDGA_TOPSHIRILGAN_SANA': xat.get('sud_topshirilgan_sana', '') or '—',
        'SUD_ISH_RAQAMI': xat.get('sud_ish_raqami', '') or '—',
        'MIB_OTKAZILGAN_SANA': xat.get('mib_otkazilgan_sana', '') or '—',
        'JILD_OCHILGAN_SANA': datetime.date.today().strftime('%d.%m.%Y'),
    }
    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_yigma_jild_titul(output_path, xat, portfel_row, mijoz, settings):
    """
    Ijro harakati (MIB) yig'ma jildi uchun titul (muqova) hujjatini,
    yuklab bo'ladigan shablon (YIGMA_JILD_TITUL_TEMPLATE_PATH) asosida
    yaratadi. Bu qog'oz holatidagi ish papkasining birinchi varag'i bo'lib,
    ish bo'yicha barcha asosiy ma'lumotlarni bir joyda ko'rsatadi.
    """
    doc = Document(_effektiv_shablon_yoli(YIGMA_JILD_TITUL_TEMPLATE_PATH, 'yigma_jild_titul_shablon.docx'))

    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '') or \
        portfel_row.get('mijoz_nomi', '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))
    mijoz_manzil = (mijoz.get('manzil') if mijoz else '') or ''
    hujjat_raqami = (mijoz.get('hujjat_raqami') if mijoz else '') or ''
    jami_qarz = (portfel_row.get('asosiy_qarz', 0) or 0) + (portfel_row.get('foiz_qarz', 0) or 0) + \
        (portfel_row.get('jarima', 0) or 0)

    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'MIB_ISH_RAQAMI': xat.get('mib_ish_raqami', '') or '—',
        'QARZDOR_ISM': mijoz_ism,
        'MIJOZ_TURI': {'yuridik': "Yuridik shaxs", 'yatt': "Yakka tartibdagi tadbirkor (YaTT)"}.get(
            xat.get('mijoz_turi'), "Jismoniy shaxs"),
        'MIJOZ_MANZIL': mijoz_manzil or '—',
        'HUJJAT_RAQAMI': hujjat_raqami or '—',
        'ANKETA_RAQAM': xat.get('anketa_raqami', '') or '—',
        'JAMI_QARZ': f"{jami_qarz:,.0f}".replace(',', ' '),
        'XAT_SANASI': _fmt_sana_iso(xat.get('yaratilgan_sana')) or '—',
        'XAT_YUBORILGAN_SANA': _fmt_sana_iso(xat.get('yuborilgan_sana')) or '—',
        'DAVO_ARIZA_SANA': _fmt_sana_iso(xat.get('davo_ariza_sana')) or '—',
        'PALATADAN_QAYTGAN_SANA': xat.get('davo_ariza_imzo_sana', '') or '—',
        'SUDGA_TOPSHIRILGAN_SANA': xat.get('sud_topshirilgan_sana', '') or '—',
        'SUD_ISH_RAQAMI': xat.get('sud_ish_raqami', '') or '—',
        'MIB_OTKAZILGAN_SANA': xat.get('mib_otkazilgan_sana', '') or '—',
        'JILD_OCHILGAN_SANA': datetime.date.today().strftime('%d.%m.%Y'),
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def _fmt_sana_iso(sana_str):
    if not sana_str:
        return ''
    try:
        return datetime.datetime.fromisoformat(sana_str).strftime('%d.%m.%Y')
    except Exception:
        return sana_str


_BIRLIK = ['', 'bir', 'ikki', 'uch', "to'rt", 'besh', 'olti', 'yetti', 'sakkiz', "to'qqiz"]
_ONLIK = ['', "o'n", 'yigirma', "o'ttiz", 'qirq', 'ellik', 'oltmish', 'yetmish', 'sakson', "to'qson"]
_DARAJA = ['', 'ming', 'million', 'milliard', 'trillion']


def _uch_xona_sozda(n):
    """0..999 oralig'idagi sonni so'z bilan yozadi."""
    qismlar = []
    if n >= 100:
        qismlar.append(f"{_BIRLIK[n // 100]} yuz" if n // 100 > 1 else 'bir yuz')
        n %= 100
    if n >= 10:
        qismlar.append(_ONLIK[n // 10])
        n %= 10
    if n > 0:
        qismlar.append(_BIRLIK[n])
    return ' '.join(qismlar)


def summa_sozda(val):
    """Summani o'zbekcha so'z bilan yozadi (rasmiy hujjatlarda summa
    qavs ichida so'z bilan takrorlanishi talab qilinadi).
    Masalan: 1 250 000 -> "bir million ikki yuz ellik ming"."""
    try:
        son = int(round(float(val or 0)))
    except (TypeError, ValueError):
        return ''
    if son == 0:
        return 'nol'
    manfiy = son < 0
    son = abs(son)
    guruhlar = []
    while son > 0:
        guruhlar.append(son % 1000)
        son //= 1000
    qismlar = []
    for i in range(len(guruhlar) - 1, -1, -1):
        if guruhlar[i] == 0:
            continue
        matn = _uch_xona_sozda(guruhlar[i])
        if i > 0:
            matn = f"{matn} {_DARAJA[i]}"
        qismlar.append(matn)
    natija = ' '.join(qismlar)
    return f"minus {natija}" if manfiy else natija


_OY_NOMLARI = ['yanvar', 'fevral', 'mart', 'aprel', 'may', 'iyun',
               'iyul', 'avgust', 'sentabr', 'oktabr', 'noyabr', 'dekabr']


def sana_rasmiy(qiymat):
    """Sanani rasmiy hujjat uslubida yozadi: 2026 yil "15" sentabr.
    Kiruvchi qiymat kk.oo.yyyy, ISO yoki date/datetime bo'lishi mumkin."""
    if not qiymat:
        return ''
    dt = None
    if isinstance(qiymat, (datetime.date, datetime.datetime)):
        dt = qiymat
    else:
        matn = str(qiymat).strip()
        for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
            try:
                dt = datetime.datetime.strptime(matn[:10], fmt)
                break
            except ValueError:
                continue
        if dt is None:
            try:
                dt = datetime.datetime.fromisoformat(matn)
            except ValueError:
                return matn  # tanib bo'lmadi — borini qaytaramiz
    return f'{dt.year} yil "{dt.day:02d}" {_OY_NOMLARI[dt.month - 1]}'


def generate_sugurta_tovon_arizasi(output_path, xat, portfel_row, undirish, settings):
    """Sug'urta kompaniyasiga yuboriladigan "sug'urta tovonini to'lash
    to'g'risida" gi ariza hujjatini tayyorlaydi.

    MUHIM: bu ariza FAQAT kredit qarzini MUDDATIDAN OLDIN undirish
    bo'yicha chiqarilgan sud qarorlari uchun beriladi — sug'urta hodisasi
    aynan shunday qaror asosida yuzaga keladi.

    xat       — xatlar jadvalidagi qator (sud/MIB ma'lumotlari bilan)
    undirish  — sugurta_undirish jadvalidagi qator (polis, ariza, summa)
    """
    doc = Document(_effektiv_shablon_yoli(SUGURTA_TOVON_TEMPLATE_PATH,
                                          'sugurta_tovon_ariza_shablon.docx'))
    undirish = undirish or {}
    portfel_row = portfel_row or {}

    asosiy_qarz = portfel_row.get('asosiy_qarz') or 0
    talab = undirish.get('talab_summasi') or asosiy_qarz

    def _qovussiz(matn):
        """Shablonda nom qo'shtirnoq ichida yozilgan — agar qiymatning
        o'zida ham qo'shtirnoq bo'lsa, ikki marta chiqib ketmasligi uchun
        chetdagilarini olib tashlaymiz."""
        return str(matn or '').strip().strip('"').strip('«»').strip()

    # Navbatdagi to'lovning muddati o'tgan sanasi: portfel kesimi sanasidan
    # DPD (kechikkan kunlar) ayirib hisoblanadi — aynan shu sanadan boshlab
    # to'lov amalga oshirilmagan.
    tolov_sanasi = ''
    try:
        dpd = int(portfel_row.get('dpd_max') or portfel_row.get('dpd_asosiy') or 0)
        if dpd > 0:
            kesim = portfel_row.get('import_sanasi')
            kesim_dt = None
            if kesim:
                for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
                    try:
                        kesim_dt = datetime.datetime.strptime(str(kesim)[:10], fmt)
                        break
                    except ValueError:
                        continue
                if kesim_dt is None:
                    try:
                        kesim_dt = datetime.datetime.fromisoformat(str(kesim))
                    except ValueError:
                        kesim_dt = None
            kesim_dt = kesim_dt or datetime.datetime.now()
            tolov_sanasi = (kesim_dt - datetime.timedelta(days=dpd)).strftime('%d.%m.%Y')
    except (TypeError, ValueError):
        pass
    if not tolov_sanasi:
        tolov_sanasi = xat.get('yuborilgan_sana') or ''

    # Kredit shartnoma raqami: portfelda alohida ustun yo'q — u anketa
    # (zayafka) raqami bilan bir xil bo'ladi. Agar xodim biznes-hamroh
    # bo'limida boshqa raqam kiritgan bo'lsa, o'sha ustun turadi.
    shartnoma_raqami = (undirish.get('kredit_shartnoma_raqami')
                        or xat.get('anketa_raqami')
                        or portfel_row.get('anketa_raqami') or '')

    mapping = {
        'SUGURTA_KOMPANIYA': _qovussiz(undirish.get('sugurta_kompaniya')),
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'BANK_MANZIL': settings.get('bank_rasmiy_manzil_filial') or settings.get('bank_manzil', ''),
        'BANK_TEL': settings.get('filial_tel') or settings.get('bank_tel', ''),

        'POLIS_SANASI': sana_rasmiy(undirish.get('polis_sanasi')),
        'SUGURTA_SHARTNOMA_RAQAM': undirish.get('polis_raqami', '') or '',
        'POLIS_RAQAM': undirish.get('polis_raqami', '') or '',

        'MIJOZ_NOMI': xat.get('mijoz_nomi', '') or portfel_row.get('mijoz_nomi', ''),
        'KREDIT_SHARTNOMA_SANASI': sana_rasmiy(portfel_row.get('shartnoma_sanasi')),
        'KREDIT_SHARTNOMA_RAQAM': shartnoma_raqami,

        'TOLOV_MUDDATI_SANASI': sana_rasmiy(tolov_sanasi),
        # Sud nomi: "muddatidan oldin undirish" da'volari iqtisodiy sudda
        # ko'riladi (Davo ariza shabloni bilan bir xil mantiq), jismoniy
        # shaxs bo'lsa — fuqarolik sudi.
        'SUD_NOMI': settings.get(
            'sud_fuqarolik_nomi' if (xat.get('mijoz_turi') == 'jismoniy'
                                     and (xat.get('davo_ariza_turi') or '').startswith('jismoniy'))
            else 'sud_iqtisodiy_nomi', '') or '',
        'QAROR_SANASI': sana_rasmiy(xat.get('sud_qaror_sana') or xat.get('sud_topshirilgan_sana')),
        'QAROR_RAQAMI': xat.get('sud_ish_raqami', '') or '',

        'ASOSIY_QARZ': _fmt_summa(asosiy_qarz),
        'ASOSIY_QARZ_SOZDA': summa_sozda(asosiy_qarz),
        'TALAB_SUMMASI': _fmt_summa(talab),

        'BANK_REKVIZIT_NOMI': f"{settings.get('bank_qisqa_nomi', '')} {settings.get('filial_nomi', '')} filiali".strip(),
        'BANK_MFO': settings.get('bank_kodi_filial') or settings.get('bh_mfo', '') or '',
        'BANK_HISOB_RAQAMI': settings.get('bank_hisob_raqami_filial', '') or '',
        'BANK_STIR': settings.get('bank_stir', '') or '',

        'BOSHQA_HUJJATLAR': undirish.get('mib_asos_hujjat_nomi', '') or '—',
        'RAHBAR_ISM': settings.get('rahbar_ism', '') or '',
        'ARIZA_SANASI': sana_rasmiy(datetime.date.today()),
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_vafot_sugurta_tovon_arizasi(output_path, vafot_row, portfel_row, undirish,
                                         mijoz, settings):
    """Sug'urtalangan shaxsning VAFOTI munosabati bilan sug'urta tovonini
    to'lash to'g'risidagi arizani tayyorlaydi.

    Tizimda mavjud ma'lumotlar (mijoz, kredit, qarz, vafot sanasi, polis)
    avtomatik qo'yiladi. Tizimda BO'LMAGAN ma'lumotlar — pasport
    seriya/raqami, FHDYo guvohnomasi raqami va sanasi, vafot sababi,
    sug'urta shartnomasining tegishli bandi — xodim tomonidan QO'LDA
    kiritiladi va sugurta_undirish yozuvida saqlanadi."""
    doc = Document(_effektiv_shablon_yoli(VAFOT_SUGURTA_TOVON_TEMPLATE_PATH,
                                          'vafot_sugurta_tovon_shablon.docx'))
    undirish = undirish or {}
    portfel_row = portfel_row or {}
    vafot_row = vafot_row or {}

    asosiy = portfel_row.get('asosiy_qarz') or 0
    foiz = portfel_row.get('foiz_qarz') or 0
    jami = asosiy + foiz
    talab = undirish.get('talab_summasi') or jami

    mijoz_nomi = vafot_row.get('mijoz_nomi') or portfel_row.get('mijoz_nomi', '')
    pinfl = portfel_row.get('pinfl') or portfel_row.get('stir') or \
        (mijoz.get('hujjat_raqami') if mijoz else '') or ''

    mapping = {
        'SUGURTA_KOMPANIYA': str(undirish.get('sugurta_kompaniya')
                                 or vafot_row.get('sugurta_kompaniya') or '').strip().strip('"'),
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'BANK_MANZIL': settings.get('bank_rasmiy_manzil_filial') or settings.get('bank_manzil', ''),
        'BANK_TEL': settings.get('filial_tel') or settings.get('bank_tel', ''),

        'ARIZA_RAQAMI': undirish.get('ariza_raqami', '') or '',
        'ARIZA_SANASI': sana_rasmiy(undirish.get('ariza_sana') or datetime.date.today()),

        # --- Sug'urta shartnomasi (qo'lda kiritiladi) ---
        'SUGURTA_SHARTNOMA_SANASI': sana_rasmiy(undirish.get('sugurta_shartnoma_sanasi')
                                                or undirish.get('polis_sanasi')),
        'SUGURTA_SHARTNOMA_RAQAMI': undirish.get('sugurta_shartnoma_raqami') or
                                    undirish.get('polis_raqami') or
                                    vafot_row.get('sugurta_polis_raqam', '') or '',
        'POLIS_RAQAM': undirish.get('polis_raqami') or vafot_row.get('sugurta_polis_raqam', '') or '',
        'SHARTNOMA_BANDI': undirish.get('shartnoma_bandi', '') or '',

        # --- Sug'urtalangan shaxs ---
        'MIJOZ_NOMI': mijoz_nomi,
        'MIJOZ_PINFL': pinfl,
        'PASPORT_SERIYA': undirish.get('pasport_seriya', '') or '',
        'PASPORT_RAQAM': undirish.get('pasport_raqam', '') or '',

        # --- Kredit ---
        'KREDIT_SHARTNOMA_SANASI': sana_rasmiy(portfel_row.get('shartnoma_sanasi')),
        'KREDIT_SHARTNOMA_RAQAM': vafot_row.get('anketa_raqami') or
                                  portfel_row.get('anketa_raqami', '') or '',
        'KREDIT_SUMMASI': _fmt_summa(portfel_row.get('jami_berilgan_summa')
                                     or portfel_row.get('ead') or 0),

        # --- Vafot (guvohnoma va sabab qo'lda kiritiladi) ---
        'VAFOT_SANASI': sana_rasmiy(vafot_row.get('vafot_sanasi')),
        'GUVOHNOMA_SANASI': sana_rasmiy(undirish.get('guvohnoma_sanasi')),
        'GUVOHNOMA_RAQAMI': undirish.get('guvohnoma_raqami', '') or '',
        'VAFOT_SABABI': undirish.get('vafot_sababi', '') or '',

        # --- Qarzdorlik ---
        'ASOSIY_QARZ': _fmt_summa(asosiy),
        'FOIZ_QARZ': _fmt_summa(foiz),
        'JAMI_QARZ': _fmt_summa(jami),
        'JAMI_QARZ_SOZDA': summa_sozda(jami),
        'TALAB_SUMMASI': _fmt_summa(talab),

        # --- Bank rekvizitlari ---
        'BANK_REKVIZIT_NOMI': f"{settings.get('bank_qisqa_nomi', '')} {settings.get('filial_nomi', '')} filiali".strip(),
        'BANK_MFO': settings.get('bank_kodi_filial') or settings.get('bh_mfo', '') or '',
        'BANK_HISOB_RAQAMI': settings.get('bank_hisob_raqami_filial', '') or '',
        'BANK_STIR': settings.get('bank_stir', '') or '',

        'BOSHQA_HUJJATLAR': undirish.get('mib_asos_hujjat_nomi', '') or '—',
        'RAHBAR_ISM': settings.get('rahbar_ism', '') or '',
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_sugurta_xabarnoma(output_path, vafot_row, portfel_row, mijoz, settings):
    """
    Vafot etgan mijozning kreditini qoplash so'rovi bilan sug'urta
    kompaniyasiga yuboriladigan xabarnoma hujjatini tayyorlaydi.
    """
    doc = Document(_effektiv_shablon_yoli(SUGURTA_XABARNOMA_TEMPLATE_PATH, 'sugurta_xabarnoma_shablon.docx'))

    jami_qarz = (portfel_row.get('asosiy_qarz', 0) or 0) + (portfel_row.get('foiz_qarz', 0) or 0) + \
        (portfel_row.get('jarima', 0) or 0)
    mijoz_ism = (mijoz.get('ism') if mijoz else None) or vafot_row.get('mijoz_nomi', '') or \
        portfel_row.get('mijoz_nomi', '')

    mapping = {
        'BANK_NOMI': settings.get('bank_nomi', ''),
        'BANK_QISQA_NOMI': settings.get('bank_qisqa_nomi', ''),
        'FILIAL_NOMI': settings.get('filial_nomi', ''),
        'RAHBAR_ISM': settings.get('rahbar_ism', ''),

        'MIJOZ_ISM': mijoz_ism,
        'MIJOZ_PINFL': portfel_row.get('pinfl', '') or (mijoz.get('hujjat_raqami') if mijoz else '') or '',
        'ANKETA_RAQAM': vafot_row.get('anketa_raqami', '') or portfel_row.get('anketa_raqami', ''),
        'SHARTNOMA_SANA': portfel_row.get('shartnoma_sanasi', '') or '',
        'KREDIT_SUMMA': _fmt_summa(portfel_row.get('jami_berilgan_summa') or portfel_row.get('ead', 0)),
        'KREDIT_TUGASH_SANASI': vafot_row.get('kredit_tugash_sanasi', '') or
                                 portfel_row.get('shartnoma_tugash_sanasi', '') or '',
        'JAMI_QARZ': _fmt_summa(jami_qarz),

        'VAFOT_SANASI': vafot_row.get('vafot_sanasi', '') or '',
        'SUGURTA_KOMPANIYA': vafot_row.get('sugurta_kompaniya', '') or '',
        'SUGURTA_POLIS_RAQAM': vafot_row.get('sugurta_polis_raqam', '') or '',
        'XABARNOMA_SANA': datetime.date.today().strftime('%d.%m.%Y'),
    }

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def _bar_chart_png(labels, values, title, color='#1E2761', fmt_short=True):
    """Oddiy gorizontal bar-chart PNG faylini vaqtinchalik joyga chizib beradi."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import tempfile

    fig, ax = plt.subplots(figsize=(7.2, max(2.2, 0.5 * len(labels) + 0.8)))
    y_pos = range(len(labels))
    ax.barh(list(y_pos), values, color=color)
    ax.set_yticks(list(y_pos))
    ax.set_yticklabels(labels, fontsize=9)
    ax.invert_yaxis()
    ax.set_title(title, fontsize=11, fontweight='bold', loc='left')
    for spine in ('top', 'right'):
        ax.spines[spine].set_visible(False)

    def _short(v):
        if fmt_short:
            if v >= 1e9:
                return f"{v/1e9:.1f} mlrd"
            if v >= 1e6:
                return f"{v/1e6:.1f} mln"
        return f"{v:,.0f}".replace(',', ' ')

    for i, v in enumerate(values):
        ax.text(v, i, ' ' + _short(v), va='center', fontsize=8.5)
    ax.set_xticks([])
    fig.tight_layout()
    tmp = tempfile.NamedTemporaryFile(suffix='.png', delete=False)
    fig.savefig(tmp.name, dpi=150)
    plt.close(fig)
    return tmp.name


def generate_tahlil_hisoboti(output_path, tahlil, settings):
    """"Tahlil" bo'limi uchun to'liq portfel tahlili hisobotini (Word) tayyorlaydi."""
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(settings.get('bank_nomi', ''))
    r.bold = True
    r.font.size = Pt(13)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f"{settings.get('filial_nomi','')} filiali")
    r2.font.size = Pt(11)

    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("PORTFEL TAHLILI HISOBOTI")
    tr.bold = True
    tr.font.size = Pt(18)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(datetime.date.today().strftime('%d.%m.%Y'))
    sr.font.size = Pt(10)
    doc.add_paragraph()

    doc.add_heading('Umumiy ko\'rsatkichlar', level=2)
    table = doc.add_table(rows=0, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr = table.add_row().cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'Ko\'rsatkich', 'Soni', 'EAD summasi (so\'m)'
    rows_data = [
        ('Jami portfel', tahlil['jami_soni'], tahlil['jami_ead']),
        ('Jismoniy shaxslar', tahlil['jismoniy']['soni'], tahlil['jismoniy']['ead']),
        ('Yuridik shaxslar', tahlil['yuridik']['soni'], tahlil['yuridik']['ead']),
    ]
    for label, soni, ead in rows_data:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = f"{soni:,}".replace(',', ' ')
        row[2].text = f"{ead:,.0f}".replace(',', ' ')

    doc.add_paragraph()
    doc.add_heading('Stage bo\'yicha taqsimot', level=2)
    table2 = doc.add_table(rows=0, cols=4)
    table2.style = 'Light Grid Accent 1'
    hdr2 = table2.add_row().cells
    hdr2[0].text, hdr2[1].text, hdr2[2].text, hdr2[3].text = 'Stage', 'Soni', 'EAD summasi', 'Ulush (%)'
    for st in tahlil['stage']:
        row = table2.add_row().cells
        row[0].text = f"Stage {st['stage']}"
        row[1].text = f"{st['soni']:,}".replace(',', ' ')
        row[2].text = f"{st['ead']:,.0f}".replace(',', ' ')
        row[3].text = f"{st['ulush']}%"

    doc.add_paragraph()
    doc.add_heading('Tarmoq (soha) kesimida', level=2)
    tarmoq_top = tahlil['tarmoq'][:10]
    labels = [t['tarmoq'][:35] for t in tarmoq_top]
    values = [t['ead'] for t in tarmoq_top]
    if labels:
        try:
            chart_path = _bar_chart_png(labels, values, "Tarmoq bo'yicha EAD summasi")
            doc.add_picture(chart_path, width=Inches(6.3))
        except ImportError:
            # matplotlib mavjud bo'lmasa, hisobot baribir jadval bilan davom etadi
            note = doc.add_paragraph()
            note.add_run("(Grafik ko'rinishi mavjud emas — matplotlib kutubxonasi topilmadi. "
                          "Jadval quyida ko'rsatilgan.)").italic = True

    table3 = doc.add_table(rows=0, cols=3)
    table3.style = 'Light Grid Accent 1'
    hdr3 = table3.add_row().cells
    hdr3[0].text, hdr3[1].text, hdr3[2].text = 'Tarmoq', 'Soni', 'EAD summasi'
    for t in tahlil['tarmoq']:
        row = table3.add_row().cells
        row[0].text = t['tarmoq']
        row[1].text = f"{t['soni']:,}".replace(',', ' ')
        row[2].text = f"{t['ead']:,.0f}".replace(',', ' ')

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_reja_hisoboti(output_path, reja, tarmoq_reja, settings):
    """"Reja Grafik" bo'limi uchun kunlik ish rejasi hisobotini (Word) tayyorlaydi."""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(settings.get('bank_nomi', ''))
    r.bold = True
    r.font.size = Pt(13)

    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("KUNLIK ISH REJASI HISOBOTI")
    tr.bold = True
    tr.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(datetime.date.today().strftime('%d.%m.%Y'))
    sr.font.size = Pt(10)
    doc.add_paragraph()

    doc.add_heading("Bugungi reja / bajarilish", level=2)
    table = doc.add_table(rows=0, cols=4)
    table.style = 'Light Grid Accent 1'
    hdr = table.add_row().cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = 'Ish turi', 'Kerak', 'Bajarildi', 'Qolib ketyapti'
    for turk in reja['turkumlar']:
        row = table.add_row().cells
        row[0].text = turk['nomi']
        row[1].text = str(turk['reja'])
        row[2].text = str(turk['bajarildi'])
        row[3].text = str(turk['qoldi'])
    doc.add_paragraph(f"Umumiy bajarilish: {reja['foiz']}%")

    doc.add_paragraph()
    doc.add_heading("Tarmoq kesimida bajarilgan ishlar", level=2)
    table2 = doc.add_table(rows=0, cols=6)
    table2.style = 'Light Grid Accent 1'
    hdr2 = table2.add_row().cells
    for i, h in enumerate(['Tarmoq', 'Jami xat', 'Yuborilgan', 'Davo ariza', 'Sudga', 'MIBga']):
        hdr2[i].text = h
    for t in tarmoq_reja:
        row = table2.add_row().cells
        row[0].text = t['tarmoq']
        row[1].text = str(t['xat_soni'])
        row[2].text = str(t['yuborilgan_soni'])
        row[3].text = str(t['davo_soni'])
        row[4].text = str(t['sud_soni'])
        row[5].text = str(t['mib_soni'])

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


# ══════════════════════════════════════════════════════════════════════
# SUDGA BERILADIGAN MA'LUMOTNOMALAR (rasmiy shakl)
# ══════════════════════════════════════════════════════════════════════
# Bu ikki hujjat sudga taqdim etiladigan RASMIY ma'lumotnomalar:
#   1) "Kredit qarzdorligi holati to'g'risida" — da'vo arizasi bilan birga
#      beriladi: shartnoma ma'lumotlari va qarzdorlikning to'liq tarkibi.
#   2) "Qo'shimcha ma'lumotnoma — qarzdorlik summasining o'zgarishi" — ish
#      sudda ko'rilayotgan paytda qarz o'zgarsa (to'lov tushsa yoki foiz
#      o'ssa) beriladi: avvalgi va bugungi holat solishtiriladi.
SUD_MALUMOTNOMA_QARZDORLIK_TEMPLATE_PATH = os.path.join(
    _base_dir(), 'templates', 'sud_malumotnoma_qarzdorlik_shablon.docx')
SUD_QOSHIMCHA_MALUMOTNOMA_TEMPLATE_PATH = os.path.join(
    _base_dir(), 'templates', 'sud_qoshimcha_malumotnoma_shablon.docx')

_OY_NOMLARI = ['yanvar', 'fevral', 'mart', 'aprel', 'may', 'iyun',
               'iyul', 'avgust', 'sentabr', 'oktabr', 'noyabr', 'dekabr']


def _sana_qismlari(sana_obj=None):
    """Hujjat sarlavhasidagi «20__ yil "__" ______» shakli uchun qismlar."""
    d = sana_obj or datetime.date.today()
    return {'YIL_2': d.strftime('%y'), 'KUN': d.strftime('%d'),
            'OY_NOMI': _OY_NOMLARI[d.month - 1]}


def _foiz_matni(qiymat):
    """Yillik foizni ortiqcha nollarsiz ko'rsatadi (24.0 -> 24)."""
    if qiymat in (None, ''):
        return '______'
    try:
        son = float(qiymat)
    except (TypeError, ValueError):
        return str(qiymat)
    return str(int(son)) if son == int(son) else str(son)


def _sud_nomini_aniqlash(xat, settings):
    """Yuridik shaxs va YaTT — iqtisodiy sudda, jismoniy shaxs — fuqarolik
    sudida ko'riladi."""
    turi = str(xat.get('mijoz_turi') or '').strip().lower()
    kalit = 'sud_iqtisodiy_nomi' if turi in ('yuridik', 'yatt', 'le') else 'sud_fuqarolik_nomi'
    return settings.get(kalit, '') or ''


def _imzo_maydonlari(settings):
    return {
        'RAHBAR_ISM': settings.get('boshqaruvchi_ism', '') or settings.get('sud_ariza_imzo_ism', '') or '____________________ (F.I.Sh.)',
        'BOSH_HISOBCHI_ISM': settings.get('bosh_hisobchi_ism', '') or '____________________ (F.I.Sh.)',
        'IJROCHI_ISM': settings.get('ijrochi_ism', '') or settings.get('bh_masul_xodim', '') or '______________________',
        'IJROCHI_TEL': settings.get('ijrochi_tel', '') or settings.get('bh_masul_tel', '') or '___________________',
    }


def _taminot_tavsifi(taminot):
    """Garov/kafillik ma'lumotini bitta matnga jamlaydi."""
    if not taminot:
        return "ta'minotsiz (blankali kredit)"
    qismlar = []
    if taminot.get('garov_tavsifi'):
        matn = str(taminot['garov_tavsifi']).strip()
        if taminot.get('garov_bahosi'):
            matn += f" (baholangan qiymati {_fmt_summa(taminot['garov_bahosi'])} so'm)"
        qismlar.append('garov: ' + matn)
    if taminot.get('kafil_ism'):
        qismlar.append('kafil: ' + str(taminot['kafil_ism']).strip())
    return '; '.join(qismlar) if qismlar else "ta'minotsiz (blankali kredit)"


def generate_sud_malumotnoma_qarzdorlik(output_path, xat, portfel_row, mijoz,
                                        settings, taminot=None, xat_raqami='',
                                        oxirgi_tolov_sanasi=None):
    """Sudga beriladigan «Kredit qarzdorligi holati to'g'risida» MA'LUMOTNOMA.

    Raqamlar MANBAI:
      • Shartnoma ma'lumotlari va berilgan kredit summasi — portfeldan;
      • Qarzdorlik tarkibi — DAVO ARIZADAGI summalardan (agar davo ariza
        tayyorlangan bo'lsa), chunki sudga aynan shu summa so'ralgan va
        ma'lumotnoma unga MOS kelishi shart. Davo ariza hali yo'q bo'lsa —
        joriy portfel qiymatlari ishlatiladi.
    """
    tpl = _effektiv_shablon_yoli(SUD_MALUMOTNOMA_QARZDORLIK_TEMPLATE_PATH,
                                 'sud_malumotnoma_qarzdorlik_shablon.docx')
    if not os.path.exists(tpl):
        raise FileNotFoundError("Shablon topilmadi: sud_malumotnoma_qarzdorlik_shablon.docx")
    doc = Document(tpl)
    portfel_row = portfel_row or {}

    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))
    turi = str(xat.get('mijoz_turi') or '').strip().lower()

    # MUHIM: hujjatdagi BARCHA raqamlar BITTA sanaga tegishli bo'lishi
    # shart — aks holda "qoldiq asosiy qarz 180 mln, shundan muddati
    # o'tgan 190 mln" kabi o'zaro zid raqamlar chiqib qoladi va sud
    # hujjatga ishonmaydi. Shuning uchun:
    #   • Davo ariza tayyorlangan bo'lsa — hamma raqam DAVO ARIZADAGI
    #     summalardan olinadi va sana ham davo ariza sanasi bo'ladi
    #     (sudga aynan shu summa so'ralgan);
    #   • Davo ariza hali yo'q bo'lsa — joriy portfel va bugungi sana.
    davo_bor = bool(xat.get('davo_summasi_asosiy') or xat.get('davo_summasi_foiz')
                    or xat.get('davo_summasi_jarima'))
    if davo_bor:
        asosiy = xat.get('davo_summasi_asosiy') or 0
        foiz = xat.get('davo_summasi_foiz') or 0
        penya = xat.get('davo_summasi_jarima') or 0
        holat_sanasi = (xat.get('sud_topshirilgan_sana') or xat.get('davo_ariza_sana')
                        or datetime.date.today())
    else:
        asosiy = portfel_row.get('asosiy_qarz') or 0
        foiz = portfel_row.get('foiz_qarz') or 0
        penya = portfel_row.get('jarima') or 0
        holat_sanasi = datetime.date.today()
    jami = asosiy + foiz + penya

    berilgan = portfel_row.get('jami_berilgan_summa') or 0
    qoldiq_asosiy = asosiy
    qaytarilgan = max(berilgan - qoldiq_asosiy, 0) if berilgan else 0

    mapping = {
        'SUD_NOMI': _sud_nomini_aniqlash(xat, settings),
        'XAT_RAQAMI': xat_raqami or '______',
        'BANK_NOMI': settings.get('bank_nomi', '') or '',
        'BANK_STIR': settings.get('bank_stir', '') or '______________',
        'BANK_MANZIL': settings.get('bank_rasmiy_manzil_filial', '') or '',
        'QARZDOR_ISM': mijoz_ism,
        'QARZDOR_TURI_IZOH': 'tashkilot nomi' if turi in ('yuridik', 'yatt', 'le') else 'F.I.Sh.',
        'PINFL_STIR': portfel_row.get('pinfl') or portfel_row.get('stir') or '—',
        'QARZDOR_MANZIL': (mijoz or {}).get('manzil') or '________________________________',
        'SHARTNOMA_SANASI': sana_rasmiy(portfel_row.get('shartnoma_sanasi')),
        'SHARTNOMA_RAQAMI': xat.get('anketa_raqami', '') or '',
        'KREDIT_SUMMA': _fmt_summa(berilgan),
        'KREDIT_SUMMA_SOZDA': summa_sozda(berilgan),
        'YILLIK_FOIZ': _foiz_matni(portfel_row.get('yillik_foiz')),
        'SHARTNOMA_TUGASH_SANASI': sana_rasmiy(portfel_row.get('shartnoma_tugash_sanasi')),
        # Bu ikkisi bank ichki hujjatida bo'ladi — qo'lda to'ldiriladi
        'TOLOV_TOPSHIRIQNOMA_SANASI': '20___ yil "___" ____________',
        'TOLOV_TOPSHIRIQNOMA_RAQAMI': '________',
        'TAMINOT_TAVSIFI': _taminot_tavsifi(taminot),
        'OXIRGI_TOLOV_SANASI': sana_rasmiy(oxirgi_tolov_sanasi) or '20___ yil "___" ____________',
        'HOLAT_SANASI': sana_rasmiy(holat_sanasi),
        'JAMI_QARZ': _fmt_summa(jami),
        'JAMI_QARZ_SOZDA': summa_sozda(jami),
        'TALABNOMA_SANASI': sana_rasmiy(xat.get('yuborilgan_sana')) or '20___ yil "___" ____________',
        'TALABNOMA_RAQAMI': '________',
        # Jadval
        'T_BERILGAN_KREDIT': _fmt_summa(berilgan),
        'T_QAYTARILGAN_ASOSIY': _fmt_summa(qaytarilgan),
        'T_QOLDIQ_ASOSIY': _fmt_summa(qoldiq_asosiy),
        'T_MUDDATI_OTGAN_ASOSIY': _fmt_summa(asosiy),
        # MUHIM: bankning hisobotida foiz qarzi YAGONA ko'rsatkich —
        # u muddati o'tgan foizni bildiradi. Shuning uchun 'hisoblangan
        # (hali muddati kelmagan)' qatoriga raqam qo'yilmaydi, aks holda
        # ustunni qo'shib chiqqan odam foizni IKKI MARTA hisoblab qo'yardi.
        'T_HISOBLANGAN_FOIZ': '\u2014',
        'T_MUDDATI_OTGAN_FOIZ': _fmt_summa(foiz),
        'T_PENYA': _fmt_summa(penya),
        'T_PENYA_KUN': str(portfel_row.get('dpd_max') or 0),
        'T_BOSHQA_TOLOV': '0',
        'T_JAMI': _fmt_summa(jami),
    }
    mapping.update(_sana_qismlari())
    mapping.update(_imzo_maydonlari(settings))
    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def _yonalish(farq):
    """Farq belgisiga qarab «kamaydi / oshdi / o'zgarmadi» so'zini beradi."""
    if farq < 0:
        return 'kamaydi'
    if farq > 0:
        return 'oshdi'
    return "o'zgarmadi"


def generate_sud_qoshimcha_malumotnoma(output_path, xat, portfel_row, mijoz, settings,
                                       tolovlar=None, xat_raqami='',
                                       avvalgi_malumotnoma_raqami='', avvalgi_sana=None):
    """Sudga beriladigan «QO'SHIMCHA MA'LUMOTNOMA — qarzdorlik summasining
    o'zgarishi to'g'risida».

    Solishtirish MANBAI:
      • «Avvalgi» ustun — DAVO ARIZADA sudga ko'rsatilgan summalar;
      • «Bugungi» ustun — portfelning ENG SO'NGGI (joriy) qiymatlari;
      • To'lovlar jadvali — shu anketaga tushgan TASDIQLANGAN to'lovlar,
        davo ariza sanasidan keyingilari.
    """
    tpl = _effektiv_shablon_yoli(SUD_QOSHIMCHA_MALUMOTNOMA_TEMPLATE_PATH,
                                 'sud_qoshimcha_malumotnoma_shablon.docx')
    if not os.path.exists(tpl):
        raise FileNotFoundError("Shablon topilmadi: sud_qoshimcha_malumotnoma_shablon.docx")
    doc = Document(tpl)
    portfel_row = portfel_row or {}

    mijoz_ism = (mijoz.get('ism') if mijoz else None) or xat.get('mijoz_nomi', '')
    mijoz_ism = util.mijoz_ism_hujjat_uchun(mijoz_ism, xat.get('mijoz_turi'))

    a_asosiy = xat.get('davo_summasi_asosiy') or 0
    a_foiz = xat.get('davo_summasi_foiz') or 0
    a_penya = xat.get('davo_summasi_jarima') or 0
    a_jami = a_asosiy + a_foiz + a_penya

    b_asosiy = portfel_row.get('asosiy_qarz') or 0
    b_foiz = portfel_row.get('foiz_qarz') or 0
    b_penya = portfel_row.get('jarima') or 0
    b_jami = b_asosiy + b_foiz + b_penya

    f_asosiy, f_foiz, f_penya = b_asosiy - a_asosiy, b_foiz - a_foiz, b_penya - a_penya
    f_jami = b_jami - a_jami

    def farq_matni(v):
        return ('+' if v > 0 else '') + _fmt_summa(v)

    avvalgi_sana_matni = sana_rasmiy(avvalgi_sana or xat.get('sud_topshirilgan_sana')
                                     or xat.get('davo_ariza_sana'))

    mapping = {
        'SUD_NOMI': _sud_nomini_aniqlash(xat, settings),
        'SUD_ISH_RAQAMI': xat.get('sud_ish_raqami', '') or '________',
        'XAT_RAQAMI': xat_raqami or '______',
        'BANK_NOMI': settings.get('bank_nomi', '') or '',
        'QARZDOR_ISM': mijoz_ism,
        'AVVALGI_MALUMOTNOMA_SANASI': avvalgi_sana_matni or '20___ yil "___" ____________',
        'AVVALGI_MALUMOTNOMA_RAQAMI': avvalgi_malumotnoma_raqami or '________',
        'AVVALGI_HOLAT_SANASI': avvalgi_sana_matni or '20___ yil "___" ____________',
        'AVVALGI_JAMI': _fmt_summa(a_jami),
        'HOLAT_SANASI': sana_rasmiy(datetime.date.today()),
        'JAMI_QARZ': _fmt_summa(b_jami),
        'JAMI_QARZ_SOZDA': summa_sozda(b_jami),
        # Solishtirish jadvali
        'A_ASOSIY': _fmt_summa(a_asosiy), 'B_ASOSIY': _fmt_summa(b_asosiy), 'F_ASOSIY': farq_matni(f_asosiy),
        # Foiz qarzi yagona ko'rsatkich (yuqoridagi izohga qarang) —
        # u faqat 'muddati o'tgan foizlar' qatorida ko'rsatiladi.
        'A_FOIZ': '\u2014', 'B_FOIZ': '\u2014', 'F_FOIZ': '\u2014',
        'A_MUDDATI_OTGAN_FOIZ': _fmt_summa(a_foiz), 'B_MUDDATI_OTGAN_FOIZ': _fmt_summa(b_foiz),
        'F_MUDDATI_OTGAN_FOIZ': farq_matni(f_foiz),
        'A_PENYA': _fmt_summa(a_penya), 'B_PENYA': _fmt_summa(b_penya), 'F_PENYA': farq_matni(f_penya),
        'A_BOSHQA': '0', 'B_BOSHQA': '0', 'F_BOSHQA': '0',
        'A_JAMI': _fmt_summa(a_jami), 'B_JAMI': _fmt_summa(b_jami), 'F_JAMI': farq_matni(f_jami),
        # Xulosa
        'FARQ_JAMI': _fmt_summa(abs(f_jami)),
        'FARQ_YONALISH': _yonalish(f_jami),
        'FARQ_ASOSIY': _fmt_summa(abs(f_asosiy)), 'FARQ_ASOSIY_YONALISH': _yonalish(f_asosiy),
        'FARQ_FOIZ': _fmt_summa(abs(f_foiz)), 'FARQ_FOIZ_YONALISH': _yonalish(f_foiz),
        'FARQ_PENYA': _fmt_summa(abs(f_penya)), 'FARQ_PENYA_YONALISH': _yonalish(f_penya),
    }
    mapping.update(_sana_qismlari())
    mapping.update(_imzo_maydonlari(settings))

    # ── To'lovlar jadvali: har bir to'lov uchun alohida qator ──────────
    tolovlar = tolovlar or []
    t2 = doc.tables[1]
    namuna_row = t2.rows[1]
    jami_row = t2.rows[-1]
    jami_tolov = 0.0
    if tolovlar:
        for i, tl in enumerate(tolovlar, start=1):
            summa = tl.get('summa') or 0
            jami_tolov += summa
            if i == 1:
                qator = namuna_row
            else:
                qator = t2.add_row()
                _row_katakcha_nusxala(namuna_row, qator)
            qiymatlar = [str(i), tl.get('sana', '') or '',
                         tl.get('tranzaksiya_raqami', '') or tl.get('izoh', '') or '',
                         _fmt_summa(summa), '', '']
            for j, q in enumerate(qiymatlar):
                if j < len(qator.cells):
                    if qator.cells[j].paragraphs:
                        p = qator.cells[j].paragraphs[0]
                        if p.runs:
                            p.runs[0].text = q
                            for r in p.runs[1:]:
                                r.text = ''
                        else:
                            p.add_run(q)
        # JAMI qatorini eng oxiriga ko'chiramiz (add_row oxiriga qo'shgani uchun)
        if len(tolovlar) > 1:
            tbl = t2._tbl
            tbl.remove(jami_row._tr)
            tbl.append(jami_row._tr)
    else:
        mapping.update({'TOLOV_NOMER': '—', 'TOLOV_SANASI': '', 'TOLOV_ASOSI': "to'lov bo'lmagan",
                        'TOLOV_ASOSIYGA': '0', 'TOLOV_FOIZGA': '0', 'TOLOV_PENYAGA': '0'})
    mapping.update({'TOLOV_JAMI_ASOSIY': _fmt_summa(jami_tolov),
                    'TOLOV_JAMI_FOIZ': '0', 'TOLOV_JAMI_PENYA': '0'})

    _replace_everywhere(doc, mapping)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


# ══════════════════════════════════════════════════════════════════════
# SHABLON BELGILARINI TEKSHIRISH
# ══════════════════════════════════════════════════════════════════════
# MUHIM (amaliyotda uchragan muammo): foydalanuvchi shablonni NOTO'G'RI
# bo'limga yuklab qo'ysa, tizim uni qabul qilardi, lekin hujjat
# yaratilganda o'sha bo'lim generatori bilmagan belgilar {{...}} ko'rinishida
# QOLIB KETARDI — va bunday, yarmi to'ldirilmagan hujjat sudga ketishi
# mumkin edi. Quyidagi funksiyalar shuning oldini oladi.

def shablon_belgilari(yol_yoki_doc):
    """Word hujjatidagi barcha {{BELGI}} nomlarini to'plam qilib qaytaradi
    (paragraflar, jadvallar, kolontitullar bilan birga)."""
    doc = yol_yoki_doc if hasattr(yol_yoki_doc, 'paragraphs') else Document(yol_yoki_doc)
    qismlar = []

    def _paragraflarni_yig(manba):
        for p in manba.paragraphs:
            qismlar.append(p.text)
        for t in manba.tables:
            for r in t.rows:
                for c in r.cells:
                    for p in c.paragraphs:
                        qismlar.append(p.text)

    _paragraflarni_yig(doc)
    for bolim in doc.sections:
        for kolontitul in (bolim.header, bolim.footer):
            try:
                _paragraflarni_yig(kolontitul)
            except Exception:
                pass
    return set(re.findall(r'\{\{([A-Za-z_0-9]+)\}\}', '\n'.join(qismlar)))


def toldirilmagan_belgilar(output_path):
    """Yaratilgan hujjatda to'ldirilmay qolgan belgilar bo'lsa — ro'yxati."""
    try:
        return sorted(shablon_belgilari(output_path))
    except Exception:
        return []


def standart_shablon_belgilari(fayl_nomi):
    """Dastur bilan kelgan (standart) shablondagi belgilar — ular aynan
    shu bo'lim generatori to'ldira oladigan belgilar hisoblanadi."""
    yol = os.path.join(_base_dir(), 'templates', fayl_nomi)
    if not os.path.exists(yol):
        return set()
    try:
        return shablon_belgilari(yol)
    except Exception:
        return set()


def shablon_qaysi_bolimniki(belgilar):
    """Berilgan belgilar to'plami QAYSI standart shablonga eng mos kelishini
    topadi — foydalanuvchiga "bu shablon aslida falon bo'limniki" deb aniq
    aytish uchun."""
    papka = os.path.join(_base_dir(), 'templates')
    if not belgilar or not os.path.isdir(papka):
        return None, 0.0
    eng_yaxshi, eng_ball = None, 0.0
    for nomi in sorted(os.listdir(papka)):
        if not nomi.endswith('.docx'):
            continue
        boshqa = standart_shablon_belgilari(nomi)
        if not boshqa:
            continue
        ball = len(belgilar & boshqa) / len(belgilar | boshqa)
        if ball > eng_ball:
            eng_yaxshi, eng_ball = nomi, ball
    return eng_yaxshi, eng_ball
