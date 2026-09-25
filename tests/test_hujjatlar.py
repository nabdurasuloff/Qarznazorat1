# -*- coding: utf-8 -*-
"""Hujjat tayyorlash testlari: shablonlar, summa so'z bilan, sanalar,
sug'urta arizalari va Excel eksportlari."""
import os
import datetime

from docx import Document


# ── Summani so'z bilan yozish ────────────────────────────────────────

def test_summa_sozda(muhit):
    letters = muhit['letters']
    tekshiruvlar = {
        0: 'nol',
        1: 'bir',
        15: "o'n besh",
        100: 'bir yuz',
        1000: 'bir ming',
        1250000: 'bir million ikki yuz ellik ming',
        98000000000: "to'qson sakkiz milliard",
    }
    for son, kutilgan in tekshiruvlar.items():
        assert letters.summa_sozda(son) == kutilgan, f"{son} -> {letters.summa_sozda(son)}"


def test_summa_sozda_katta_son(muhit):
    letters = muhit['letters']
    natija = letters.summa_sozda(143251590273)
    assert natija.startswith('bir yuz qirq uch milliard')
    assert 'million' in natija and 'ming' in natija


def test_sana_rasmiy(muhit):
    letters = muhit['letters']
    assert letters.sana_rasmiy('15.09.2026') == '2026 yil "15" sentabr'
    assert letters.sana_rasmiy('2026-07-23') == '2026 yil "23" iyul'
    assert letters.sana_rasmiy('') == ''
    assert letters.sana_rasmiy(None) == ''


# ── Shablonlar joyida va to'liqmi ────────────────────────────────────

def test_shablonlar_mavjud(muhit):
    letters = muhit['letters']
    kerakli = [
        letters.SUGURTA_TOVON_TEMPLATE_PATH,
        letters.VAFOT_SUGURTA_TOVON_TEMPLATE_PATH,
        letters.TEMPLATE_PATH,
    ]
    for yol in kerakli:
        assert os.path.exists(yol), f"Shablon topilmadi: {yol}"


def test_shablonlarda_joy_tutuvchilar_bor(muhit):
    letters = muhit['letters']
    for yol in (letters.SUGURTA_TOVON_TEMPLATE_PATH,
                letters.VAFOT_SUGURTA_TOVON_TEMPLATE_PATH):
        matn = '\n'.join(p.text for p in Document(yol).paragraphs)
        assert '{{' in matn, f"{os.path.basename(yol)} da joy tutuvchi yo'q"


# ── Sug'urta arizalari tayyorlanishi ─────────────────────────────────

def test_mib_sugurta_arizasi_toldiriladi(muhit, mijoz):
    db, letters = muhit['db'], muhit['letters']
    xat = {'id': 1, 'anketa_raqami': '10001', 'mijoz_nomi': mijoz['mijoz_nomi'],
           'sud_qaror_sana': '18.06.2026', 'sud_ish_raqami': '1601-26/2200',
           'mijoz_turi': 'jismoniy', 'davo_ariza_turi': 'muddatidan_oldin'}
    undirish = {'sugurta_kompaniya': 'TEST SUGURTA', 'polis_raqami': 'POL-123',
                'polis_sanasi': '12.03.2025', 'talab_summasi': 55000000,
                'mib_asos_hujjat_nomi': 'Dalolatnoma'}
    chiqish = os.path.join(muhit['papka'], 'ariza.docx')
    letters.generate_sugurta_tovon_arizasi(chiqish, xat, mijoz, undirish,
                                           db.get_all_settings())
    assert os.path.exists(chiqish)
    matn = '\n'.join(p.text for p in Document(chiqish).paragraphs)
    assert '{{' not in matn, "Joy tutuvchi to'ldirilmay qolgan"
    assert 'TEST SUGURTA' in matn
    assert 'POL-123' in matn
    assert '1601-26/2200' in matn
    # Summa so'z bilan ham yozilishi kerak
    assert 'million' in matn


def test_vafot_sugurta_arizasi_toldiriladi(muhit, mijoz):
    db, letters = muhit['db'], muhit['letters']
    vafot = {'id': 1, 'anketa_raqami': '10001', 'mijoz_nomi': mijoz['mijoz_nomi'],
             'vafot_sanasi': '15.08.2026', 'sugurta_kompaniya': 'KAFOLAT'}
    undirish = {'sugurta_kompaniya': 'KAFOLAT', 'polis_raqami': 'HF-55',
                'guvohnoma_raqami': 'II-SR 447120', 'guvohnoma_sanasi': '20.08.2026',
                'vafot_sababi': 'yurak xuruji', 'shartnoma_bandi': '4.2',
                'pasport_seriya': 'AD', 'pasport_raqam': '2287451'}
    chiqish = os.path.join(muhit['papka'], 'vafot_ariza.docx')
    letters.generate_vafot_sugurta_tovon_arizasi(chiqish, vafot, mijoz, undirish,
                                                 None, db.get_all_settings())
    assert os.path.exists(chiqish)
    matn = '\n'.join(p.text for p in Document(chiqish).paragraphs)
    assert '{{' not in matn, "Joy tutuvchi to'ldirilmay qolgan"
    for kutilgan in ('KAFOLAT', 'II-SR 447120', 'yurak xuruji', 'AD', '2287451'):
        assert kutilgan in matn, f"{kutilgan} arizada yo'q"
    # Asosiy qarz + foiz = jami bo'lishi kerak
    assert '50 000 000' in matn and '5 000 000' in matn and '55 000 000' in matn


def test_ariza_kompaniyasiz_tayyorlanmaydi(muhit, mijoz, klient):
    """Sug'urta kompaniyasi ko'rsatilmasa — ariza tayyorlanmasligi kerak."""
    db = muhit['db']
    conn = db.get_conn()
    conn.execute('''INSERT INTO xatlar (portfel_id, anketa_raqami, mijoz_nomi,
                    mib_holati, davo_ariza_turi, yaratilgan_sana)
                    VALUES (?, '10001', ?, 'otkazildi', 'muddatidan_oldin', ?)''',
                 (mijoz['id'], mijoz['mijoz_nomi'], datetime.datetime.now().isoformat()))
    conn.commit()
    conn.close()
    r = klient.post('/api/sugurta_undirish/ariza_yaratish', json={'anketa_raqami': '10001'})
    assert r.status_code == 400
    assert 'kompaniya' in r.get_json()['xato']


# ── Holat o'zgartirish shartlari ─────────────────────────────────────

def test_asos_hujjatsiz_ariza_yuborilmaydi(muhit, mijoz, klient):
    db = muhit['db']
    conn = db.get_conn()
    conn.execute('''INSERT INTO xatlar (portfel_id, anketa_raqami, mijoz_nomi,
                    mib_holati, davo_ariza_turi, yaratilgan_sana)
                    VALUES (?, '10001', ?, 'otkazildi', 'muddatidan_oldin', ?)''',
                 (mijoz['id'], mijoz['mijoz_nomi'], datetime.datetime.now().isoformat()))
    conn.commit()
    conn.close()
    r = klient.post('/api/sugurta_undirish/holat',
                    data={'anketa_raqami': '10001', 'ish_turi': 'mib',
                          'holati': 'yuborildi', 'sugurta_kompaniya': 'TEST'})
    assert r.status_code == 400
    assert 'ASOS' in r.get_json()['xato']


def test_guvohnomasiz_vafot_arizasi_yuborilmaydi(muhit, mijoz, klient):
    db = muhit['db']
    db.mark_vafot_etgan(mijoz['id'], '10001', mijoz['mijoz_nomi'],
                        '15.08.2026', None, None)   # guvohnoma YO'Q
    r = klient.post('/api/sugurta_undirish/holat',
                    data={'anketa_raqami': '10001', 'ish_turi': 'vafot',
                          'holati': 'yuborildi', 'sugurta_kompaniya': 'TEST'})
    assert r.status_code == 400
    assert 'guvohnoma' in r.get_json()['xato'].lower()


def test_tolangan_summasiz_tolandi_belgilanmaydi(muhit, mijoz, klient):
    db = muhit['db']
    db.mark_vafot_etgan(mijoz['id'], '10001', mijoz['mijoz_nomi'],
                        '15.08.2026', None, None)
    r = klient.post('/api/sugurta_undirish/holat',
                    data={'anketa_raqami': '10001', 'ish_turi': 'vafot',
                          'holati': 'tolandi'})
    assert r.status_code == 400


def test_sababsiz_rad_etilmaydi(muhit, mijoz, klient):
    db = muhit['db']
    db.mark_vafot_etgan(mijoz['id'], '10001', mijoz['mijoz_nomi'],
                        '15.08.2026', None, None)
    r = klient.post('/api/sugurta_undirish/holat',
                    data={'anketa_raqami': '10001', 'ish_turi': 'vafot',
                          'holati': 'rad_etildi'})
    assert r.status_code == 400
