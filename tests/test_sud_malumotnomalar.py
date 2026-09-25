# -*- coding: utf-8 -*-
"""SUDGA BERILADIGAN RASMIY MA'LUMOTNOMALAR.

Ikki hujjat:
  1) «MA'LUMOTNOMA — kredit qarzdorligi holati to'g'risida» — da'vo
     arizasi bilan birga sudga beriladi;
  2) «QO'SHIMCHA MA'LUMOTNOMA — qarzdorlik summasining o'zgarishi
     to'g'risida» — ish sudda ko'rilayotganda qarz o'zgarsa beriladi.

Eng muhim talab: hujjat ICHIDAGI raqamlar o'zaro ZID BO'LMASLIGI kerak —
sud zid raqamli hujjatni qabul qilmaydi.
"""
import datetime
import os

import pytest
from docx import Document


def _matn(yol):
    d = Document(yol)
    qismlar = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            qismlar.append(' | '.join(c.text for c in r.cells))
    return '\n'.join(qismlar)


def _son(matn):
    """«190 000 000» -> 190000000"""
    toza = matn.replace(' ', ' ').replace(' ', '').replace('+', '')
    try:
        return float(toza)
    except ValueError:
        return None


@pytest.fixture()
def ish(muhit, mijoz):
    """Sudga topshirilgan, davo arizasi tayyor ish."""
    db = muhit['db']
    papka = os.path.join(muhit['papka'], 'jild')
    os.makedirs(papka, exist_ok=True)
    for kalit, qiymat in [('bank_nomi', 'AGROBANK'), ('boshqaruvchi_ism', 'A.RAXIMOV'),
                          ('bosh_hisobchi_ism', 'M.KARIMOVA'), ('ijrochi_ism', 'N.TOSHEV'),
                          ('ijrochi_tel', '+998 67 123-45-67')]:
        db.set_setting(kalit, qiymat)
    conn = db.get_conn()
    conn.execute('''UPDATE portfel SET jami_berilgan_summa=250000000, yillik_foiz=24,
                    asosiy_qarz=180000000, foiz_qarz=24500000, jarima=6200000,
                    shartnoma_sanasi='12.03.2024', shartnoma_tugash_sanasi='12.03.2027'
                    WHERE id=?''', (mijoz['id'],))
    conn.execute('''INSERT INTO xatlar (portfel_id, anketa_raqami, mijoz_nomi, mijoz_turi,
        xat_turi, holat, yaratilgan_sana, yuborilgan_sana, fayl_yoli, davo_ariza_holati,
        davo_ariza_turi, davo_ariza_fayl_yoli, davo_ariza_sana,
        davo_summasi_asosiy, davo_summasi_foiz, davo_summasi_jarima,
        sud_holati, sud_topshirilgan_sana, sud_ish_raqami, sud_yigma_jild_papka)
        VALUES (?, '10001', ?, 'yuridik', 'talabnoma_yuridik', 'yuborildi', ?, '01.08.2026',
                '/tmp/x.docx', 'olib_kelindi', 'oddiy', '/tmp/d.docx', '20.08.2026',
                190000000, 22000000, 5000000, 'topshirildi', '25.08.2026', 'A-777', ?)''',
        (mijoz['id'], mijoz['mijoz_nomi'], datetime.datetime.now().isoformat(), papka))
    conn.execute('''INSERT INTO davo_taminot (anketa_raqami, taminot_turi, kafil_ism,
        garov_tavsifi, garov_bahosi) VALUES ('10001', 'kafil_garov', 'ASROROV BOBUR',
        'Ishlab chiqarish binosi', 300000000)''')
    for sana, summa in [('12.07.2026', 5000000), ('05.09.2026', 8000000), ('18.09.2026', 12000000)]:
        conn.execute('''INSERT INTO tolovlar (manba, sana, summa, anketa_raqami, mijoz_nomi,
            holati, tranzaksiya_raqami) VALUES ('mib', ?, ?, '10001', ?, 'tasdiqlangan', ?)''',
            (sana, summa, mijoz['mijoz_nomi'], 'TR-' + sana.replace('.', '')))
    conn.commit()
    conn.close()
    return {'papka': papka}


# ── Yaratish ─────────────────────────────────────────────────────────

def test_ikkala_malumotnoma_yaratiladi(muhit, ish, klient):
    for turi, fayl in (('qarzdorlik', '06_Malumotnoma_kredit_qarzdorlik.docx'),
                       ('qoshimcha', '07_Qoshimcha_malumotnoma_ozgarish.docx')):
        r = klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                        json={'anketa_raqami': '10001', 'turi': turi})
        assert r.status_code == 200, r.get_json()
        assert os.path.exists(os.path.join(ish['papka'], fayl))


def test_notogri_tur_rad_etiladi(muhit, ish, klient):
    r = klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                    json={'anketa_raqami': '10001', 'turi': 'yoq'})
    assert r.status_code == 400


def test_jildsiz_yaratilmaydi(muhit, mijoz, klient):
    conn = muhit['db'].get_conn()
    conn.execute('''INSERT INTO xatlar (portfel_id, anketa_raqami, mijoz_nomi, mijoz_turi,
        xat_turi, holat) VALUES (?, '10001', ?, 'yuridik', 'talabnoma', 'yuborildi')''',
        (mijoz['id'], mijoz['mijoz_nomi']))
    conn.commit()
    conn.close()
    r = klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                    json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    assert r.status_code == 400
    assert "jild" in r.get_json()['xato'].lower()


# ── Hujjat ichidagi raqamlar zid bo'lmasligi ─────────────────────────

def test_qarzdorlik_malumotnomasi_ichki_ziddiyatsiz(muhit, ish, klient):
    """MUHIM: «qoldiq asosiy qarz» «shundan muddati o'tgan» dan KICHIK
    bo'lib qolmasligi kerak, va JAMI = qatorlar yig'indisiga teng bo'lishi
    shart — aks holda sud hujjatni qaytaradi."""
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    d = Document(os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx'))
    t = d.tables[0]
    qoldiq = _son(t.rows[3].cells[2].text)
    muddati_otgan = _son(t.rows[3].cells[3].text.split(':')[-1])
    penya = _son(t.rows[6].cells[2].text)
    m_foiz = _son(t.rows[5].cells[2].text)
    jami = _son(t.rows[8].cells[2].text)

    assert muddati_otgan <= qoldiq, \
        f"Muddati o'tgan ({muddati_otgan}) qoldiqdan ({qoldiq}) katta — zid ma'lumot"
    assert jami == qoldiq + m_foiz + penya, \
        f"JAMI ({jami}) qatorlar yig'indisiga ({qoldiq + m_foiz + penya}) teng emas"


def test_foiz_ikki_marta_hisoblanmaydi(muhit, ish, klient):
    """Foiz qarzi bir ustunda IKKI qatorda takrorlansa, uni qo'shib
    chiqqan odam summani ikki marta hisoblab yuborardi."""
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    t = Document(os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx')).tables[0]
    hisoblangan = t.rows[4].cells[2].text.strip()
    assert _son(hisoblangan) in (None, 0), \
        f"'Hisoblangan foizlar' qatorida ham raqam bor ({hisoblangan}) — foiz ikki marta sanaladi"


def test_oxirgi_tolov_hujjat_sanasidan_keyin_emas(muhit, ish, klient):
    """«25 avgust holatiga» deb yozilgan hujjatda «oxirgi to'lov 18
    sentabrda» degan zid jumla chiqmasligi kerak."""
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    matn = _matn(os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx'))
    assert 'oxirgi to‘lov 2026 yil "12" iyul' in matn, \
        "Oxirgi to'lov sanasi hujjat sanasiga mos emas:\n" + matn[:900]


def test_qoshimcha_malumotnoma_farqlari_togri(muhit, ish, klient):
    """Avvalgi (da'vodagi) va bugungi summa solishtiriladi, farq to'g'ri
    hisoblanadi: 217 000 000 -> 210 700 000, farq -6 300 000."""
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qoshimcha'})
    d = Document(os.path.join(ish['papka'], '07_Qoshimcha_malumotnoma_ozgarish.docx'))
    jami_row = d.tables[0].rows[6]
    avvalgi = _son(jami_row.cells[2].text)
    bugungi = _son(jami_row.cells[3].text)
    farq = _son(jami_row.cells[4].text)
    assert avvalgi == 217000000
    assert bugungi == 210700000
    assert abs(farq) == 6300000
    assert '-' in jami_row.cells[4].text, "Kamayish minus bilan ko'rsatilmagan"

    matn = _matn(os.path.join(ish['papka'], '07_Qoshimcha_malumotnoma_ozgarish.docx'))
    assert 'kamaydi' in matn, "Xulosada yo'nalish (kamaydi/oshdi) ko'rsatilmagan"


def test_qoshimchada_faqat_keyingi_tolovlar(muhit, ish, klient):
    """Jadvalga faqat sudga topshirilgandan KEYIN tushgan to'lovlar
    kirishi kerak (12.07 — oldin, shuning uchun kirmaydi)."""
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qoshimcha'})
    d = Document(os.path.join(ish['papka'], '07_Qoshimcha_malumotnoma_ozgarish.docx'))
    tolov_matni = '\n'.join(' | '.join(c.text for c in r.cells) for r in d.tables[1].rows)
    assert '05.09.2026' in tolov_matni and '18.09.2026' in tolov_matni
    assert '12.07.2026' not in tolov_matni, "Sudga topshirishdan OLDINGI to'lov jadvalga kirgan"
    assert _son(d.tables[1].rows[-1].cells[3].text) == 20000000


# ── To'ldirilmagan belgilar qolmasligi ───────────────────────────────

def test_belgilar_toliq_almashtirilgan(muhit, ish, klient):
    """Hujjatda {{...}} ko'rinishidagi to'ldirilmagan belgi qolmasligi kerak."""
    for turi, fayl in (('qarzdorlik', '06_Malumotnoma_kredit_qarzdorlik.docx'),
                       ('qoshimcha', '07_Qoshimcha_malumotnoma_ozgarish.docx')):
        klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                    json={'anketa_raqami': '10001', 'turi': turi})
        matn = _matn(os.path.join(ish['papka'], fayl))
        assert '{{' not in matn, f"{fayl} da to'ldirilmagan belgi qoldi:\n" + \
            '\n'.join(q for q in matn.split('\n') if '{{' in q)


def test_mijoz_va_bank_malumotlari_bor(muhit, ish, klient):
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    matn = _matn(os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx'))
    for kerak in ('AGROBANK', 'TESTOV', 'A.RAXIMOV', 'M.KARIMOVA', 'N.TOSHEV',
                  'ASROROV BOBUR', '250 000 000'):
        assert kerak in matn, f"'{kerak}' hujjatda yo'q"


def test_yuridik_shaxsga_iqtisodiy_sud(muhit, ish, klient):
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    matn = _matn(os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx'))
    assert 'iqtisodiy sudi' in matn, "Yuridik shaxs ishi iqtisodiy sudga yo'naltirilmagan"


# ── Shablonlar va yig'ma jild ────────────────────────────────────────

def test_shablonlar_yuklab_olinadi(muhit, klient):
    """Sozlamalarda ikkala shablonni ko'rish/yangilash mumkin bo'lishi kerak."""
    for turi in ('sud_malumotnoma_qarzdorlik', 'sud_qoshimcha_malumotnoma'):
        r = klient.get(f'/api/shablon/korish?turi={turi}')
        assert r.status_code == 200, f"{turi} shabloni ochilmadi"
        assert len(r.data) > 5000


def test_malumotnomalar_yigma_jildga_qoshiladi(muhit, ish, klient):
    """Tayyor PDF jildda ikkala ma'lumotnoma ham bo'lishi kerak."""
    holat = klient.get('/api/sud/yigma_jild_holati?anketa=10001').get_json()
    maydonlar = {h['maydon'] for h in holat['hujjatlar']}
    assert 'sud_malumotnoma_qarzdorlik_fayl' in maydonlar
    assert 'sud_qoshimcha_malumotnoma_fayl' in maydonlar
    # Bu hujjatlar MAJBURIY emas — ularsiz ham sudga ruxsat berilishi kerak
    for h in holat['hujjatlar']:
        if h['maydon'].startswith('sud_malumotnoma_qarzdorlik') or \
           h['maydon'].startswith('sud_qoshimcha_malumotnoma'):
            assert h.get('majburiy') is False
            assert h.get('tizim_yaratadi')


# ── Shablon NOTO'G'RI bo'limga yuklanishi ────────────────────────────
# AMALIYOTDA UCHRAGAN XATO: foydalanuvchi yangi shablonni eski
# "Ma'lumotnoma (sudga topshirishda)" qatoriga yukladi. Tizim qabul
# qildi, lekin hujjat yaratilganda belgilarning ko'pi {{...}} bo'lib
# qolib ketdi — chunki o'sha eski bo'lim generatori ularni bilmaydi.

def _shablon_baytlari(nomi):
    import os as _os
    yol = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
                        'backend', 'templates', nomi)
    return open(yol, 'rb').read()


def test_notogri_bolimga_shablon_yuklanmaydi(muhit, klient):
    import io
    r = klient.post('/api/shablon/yuklash', data={
        'turi': 'malumotnoma_topshirishda',
        'file': (io.BytesIO(_shablon_baytlari('sud_malumotnoma_qarzdorlik_shablon.docx')),
                 'shablon.docx')}, content_type='multipart/form-data')
    assert r.status_code == 400, "Noto'g'ri bo'limga yuklangan shablon qabul qilindi!"
    xato = r.get_json()['xato']
    assert 'mos kelmaydi' in xato
    # Qaysi bo'limga yuklash kerakligi AYTILISHI kerak
    assert 'kredit qarzdorligi holati' in xato, f"To'g'ri bo'lim ko'rsatilmagan:\n{xato}"


def test_togri_bolimga_shablon_yuklanadi(muhit, klient):
    import io
    for turi, nomi in (('sud_malumotnoma_qarzdorlik', 'sud_malumotnoma_qarzdorlik_shablon.docx'),
                       ('sud_qoshimcha_malumotnoma', 'sud_qoshimcha_malumotnoma_shablon.docx')):
        r = klient.post('/api/shablon/yuklash', data={
            'turi': turi, 'file': (io.BytesIO(_shablon_baytlari(nomi)), 'shablon.docx')},
            content_type='multipart/form-data')
        assert r.status_code == 200, f"{turi}: {r.get_json()}"


def test_ikki_yangi_shablon_ozaro_almashmaydi(muhit, klient):
    """Ikkala yangi shablon bir-birining o'rniga ham yuklanmasligi kerak."""
    import io
    r = klient.post('/api/shablon/yuklash', data={
        'turi': 'sud_qoshimcha_malumotnoma',
        'file': (io.BytesIO(_shablon_baytlari('sud_malumotnoma_qarzdorlik_shablon.docx')),
                 'shablon.docx')}, content_type='multipart/form-data')
    assert r.status_code == 400


def test_toldirilmagan_belgi_aniqlanadi(muhit, ish, klient):
    """Agar shablon baribir mos kelmasa, hujjat yaratilgandan keyin
    foydalanuvchi ogohlantirish olishi kerak."""
    letters = muhit['letters']
    # Yaratilgan normal hujjatda belgi qolmasligi kerak
    klient.post('/api/sud/rasmiy_malumotnoma_yaratish',
                json={'anketa_raqami': '10001', 'turi': 'qarzdorlik'})
    import os as _os
    qolgan = letters.toldirilmagan_belgilar(
        _os.path.join(ish['papka'], '06_Malumotnoma_kredit_qarzdorlik.docx'))
    assert qolgan == [], f"To'ldirilmagan belgilar qoldi: {qolgan}"

    # Tekshiruvchi funksiya haqiqatan ham belgini ko'ra oladimi?
    belgilar = letters.shablon_belgilari(
        _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))),
                      'backend', 'templates', 'sud_malumotnoma_qarzdorlik_shablon.docx'))
    assert 'SUD_NOMI' in belgilar and 'T_JAMI' in belgilar
